"""AdigiPQR curator: PAS-X raw export -> QA's APQR workbook, every value from an approved rule.

    python apqr_curate.py EXPORT.xlsx [EXPORT2.csv ...] --template "APQR blank template.xlsx" [-o OUT.xlsx]
                          [--batches 2605012278,...] [--validate "QA's own workbook.xlsx"]
    python apqr_curate.py --export-rules rules_F1M00011.csv          (the rule registry, for QA review)
    python apqr_curate.py EXPORT.xlsx --template T.xlsx --rules rules_F1M00011.csv   (QA's edited registry)
    python apqr_curate.py EXPORT.xlsx --learn "QA's workbook of the product.xlsx"      (new product: learn QA's layout)

Input   the export of final/APQR_FINAL_QUERY.sql (v6.2: every repetition of a step, formula / text / list activities).
Output  the template workbook exactly as it is (headers, sample rows, colours, formulas), the batches added below its
        last batch in manufacturing order, and three sheets:
          APQR_Source_Trace  one row per cell: rule, PAS-X step code path, parameter code, readings, calculation
          APQR_Exceptions    every cell left blank because a rule could not be met (missing side, units, disagreement)
          APQR_Rules         the rule set used (rule id, version, status, evidence)
        Without --template the built-in copy of the same headers is used (QA's F1M00011 workbook, header rows only).

Principles (QA 30-Oct-2026):
  - A rule names the exact STEP_CODE_PATH and PARAMETER_CODE; display names are never matched. A text / formula
    activity whose code is not yet known is named by its exact PAS-X description (status PROVISIONAL).
  - Min = lowest reading over every repetition of the named steps, Max = highest. Two sides (LHS / RHS) are combined
    only when both are present. Granulation drying (Inlet / Product / Exhaust / Flap): the start-up reading of
    'Drying Inprocess' is left out, as QA does (2605012278: Lot-I inlet 24 | 36 | 36 | 37 | 36 | 37 -> 36 / 37).
  - Coating Min / Max = PAS-X's own 'Spraying Min & Max' record, cross-checked against the 15-minute readings; a
    disagreement leaves the cell blank with an exception.
  - Nothing is invented: no constants, no guessed sources. Blank + reason beats a plausible wrong value.
  - No Ave (average) values are written (QA decision 30-Oct-2026); QA's Ave columns stay blank.
  - An MBR without an approved rule set: the generic QA format (QA's sheets and header levels, filled only on exact code
    patterns), or (--qa / --learn) QA's own layout with
    rules learned from QA's values for batches both files have. Marked FILLED_AUTO / FILLED_LEARNED: QA to approve.
Needs   Python 3.8+ and openpyxl (pip install openpyxl). This file is self-contained.
Same file as adigipqr/curate.py (portal backend); tests/test_curate.py keeps them identical.
Prototype for QA review: verify against the batch record before use in a GMP APQR."""
import argparse
import bisect
import copy
import csv
import io
import os
import re
import sys
from collections import OrderedDict

VERSION = '2.3.0'
TRACE_SHEET, EXC_SHEET, RULE_SHEET, VALID_SHEET = 'APQR_Source_Trace', 'APQR_Exceptions', 'APQR_Rules', 'APQR_Validation'
# every column of the export the curator reads; tests/test_sql.py checks each one is an output column of the SQL
EXPORT_FIELDS = ('BATCH_NO', 'MATERIAL_CODE', 'PRODUCT_NAME', 'MANUFACTURING_ORDER', 'MBR_ID', 'MBR_VERSION', 'BATCH_SIZE',
                 'BATCH_SIZE_UOM', 'MFG_START', 'MFG_MONTH', 'EXP_MONTH', 'SHELF_LIFE', 'APQR_SHEET', 'STAGE', 'STEP',
                 'STEP_PATH', 'STEP_CODE_PATH', 'APQR_PARAMETER', 'PARAMETER_CODE', 'UOM', 'SPEC_MIN', 'SPEC_MAX',
                 'TARGET_SET_VALUE', 'READING_MIN', 'READING_MAX', 'READING_AVG', 'READING_COUNT',
                 'READING_MIN_AFTER_FIRST', 'READING_MAX_AFTER_FIRST', 'READING_VALUES', 'INDIVIDUAL_MIN',
                 'INDIVIDUAL_MAX', 'INDIVIDUAL_AVG', 'INDIVIDUAL_COUNT', 'CORRECTED_ENTRIES', 'SUPERSEDED_VALUES',
                 'ACTUAL_TEXT', 'FIRST_READING_ON', 'LAST_READING_ON')
LEARN_SHEET, UNMAPPED_SHEET = 'APQR_Learning', 'APQR_Unmapped_PASX'
OWN_SHEETS = (TRACE_SHEET, EXC_SHEET, RULE_SHEET, VALID_SHEET, LEARN_SHEET, UNMAPPED_SHEET, 'AdigiPQR source')


# ------------------------------------------------------------------------------------------------ reading
def s(v):
    return '' if v is None else str(v).strip()


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if v != v else float(v)
    t = s(v).replace(',', '')
    return float(t) if re.fullmatch(r'[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?', t) else None


def cell_text(v):
    """Export cell -> text as PL/SQL Developer shows it: Excel gives 30000838.0 for 30000838, datetimes for dates."""
    import datetime as _dt
    if isinstance(v, bool) or v is None:
        return v
    if isinstance(v, int) or (isinstance(v, float) and v == v and v.is_integer() and abs(v) < 1e15):
        return str(int(v))                             # as the CSV export: 30000838, not 30000838.0
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime('%d-%m-%Y %H:%M') if isinstance(v, _dt.datetime) else v.strftime('%d-%m-%Y')
    return v


def read_export(path_or_bytes, name=None):
    """Rows (dicts) of a CSV or .xlsx export. The header is the first row (of the first 30) with BATCH_NO and
    STEP_CODE_PATH, on any sheet.
    QA's APQR workbook given by mistake is named as such (ValueError)."""
    name = name or (path_or_bytes if isinstance(path_or_bytes, str) else 'export.csv')
    data = open(path_or_bytes, 'rb').read() if isinstance(path_or_bytes, str) else path_or_bytes
    rows, hdr = [], None
    if name.lower().endswith(('.xlsx', '.xlsm')):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if hdr is None:
                    if i >= 30:
                        break
                    if {'BATCH_NO', 'STEP_CODE_PATH'} <= {s(x).upper() for x in r}:     # the export, not a trace sheet
                        hdr = [s(x).upper() for x in r]
                    continue
                if any(x is not None for x in r):
                    rows.append({h: cell_text(x) for h, x in zip(hdr, r) if h})
            if hdr:
                break
        names = wb.sheetnames
        wb.close()
        if hdr is None and any(n.strip().upper() in ('BMR DETAILS', 'BMR DETAIL', TRACE_SHEET.upper(), 'COMPRESSION')
                               for n in names):
            raise ValueError('%s is an APQR workbook (sheets %s ...), not the PAS-X export: give the export of '
                             'APQR_FINAL_QUERY.sql; QA\'s workbook goes with --qa' % (os.path.basename(name), ', '.join(names[:4])))
        return rows
    text = data.decode('utf-8-sig', errors='replace')
    delim = max(',;\t|', key=text.split('\n', 1)[0].count)
    for r in csv.reader(io.StringIO(text), delimiter=delim):
        if hdr is None:
            if any(x.strip().strip('"').upper() == 'BATCH_NO' for x in r):
                hdr = [x.strip().strip('"').upper() for x in r]
            continue
        if any(x.strip() for x in r):
            rows.append({h: re.sub(r'^(\d+)\.0+$', r'\1', x.strip()) for h, x in zip(hdr, r)})
    return rows


def batch_no(v):
    return re.sub(r'\.0$', '', s(v))


def code(v):
    """'GRAN >  LOT1 > IPC ' -> 'GRAN > LOT1 > IPC' (spacing as PAS-X writes it)."""
    return ' > '.join(p.strip() for p in s(v).split('>'))


def by_batch(rows):
    b = OrderedDict()
    for r in rows:
        b.setdefault(batch_no(r.get('BATCH_NO')), []).append(r)
    return b


def start_key(t):
    """'07-08-2026 16:33' -> '2026-08-07 16:33' (sortable)."""
    m = re.match(r'(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})(?:\D+(\d{1,2}):(\d{2}))?', s(t))
    if m:
        return '%s-%02d-%02d %02d:%s' % (m[3], int(m[2]), int(m[1]), int(m[4] or 0), m[5] or '00')
    return s(t) or '9999'


def first(b, field):
    return next((s(r.get(field)) for r in b if s(r.get(field))), '')


def stage_kg(b):
    for r in b:
        k = re.search(r'\(\s*([\d.,]+)\s*kgs?\s*\)', s(r.get('STAGE')), re.I)
        if k:
            return num(k.group(1))
    return None


def clean(v):
    """Float noise off (46.7999999999999972 -> 46.8); whole numbers as int."""
    if isinstance(v, float):
        v = round(float('%.10g' % v), 4)
        return int(v) if v.is_integer() else v
    return v


# ------------------------------------------------------------------------------------------------ rule registry
# A rule fills one QA column of one sheet. KIND:
#   PASX      value from the export: SOURCES = 'STEP_CODE_PATH|PARAMETER_CODE' (exact) separated by ';'.
#             'PATH > *|...' = any step below PATH;  '|@Description' = exact PAS-X activity description (used for
#             formula / text activities whose code is not yet confirmed - status PROVISIONAL).
#   HEADER    batch header field of the export (BATCH_SIZE, MFG_MONTH, EXP_MONTH, BMR_NO, MATERIAL_CODE).
#   FORMULA   QA's own Excel formula, written as QA writes it ({r} = row).
#   QA_ENTRY / PRINTOUT / NOT_IN_MBR / NO_PASX_SOURCE   nothing is written; the reason is traced.
# STATISTIC  MIN / MAX / AVG (count-weighted) / VALUE (one value) / SUM / MIN_AFTER_FIRST / MAX_AFTER_FIRST /
#            TEXT (entered texts) / IMIN / IMAX / IAVG (individual units).
# REQUIRE    ALL = every source must be present (both sides), else blank + exception; ANY = whatever is present.
# MACHINE    73 / 75: the column block of one compression machine; a batch made on the other machine gets NA.
# ZERO_AS    text written instead of 0 (QA writes NA for no rejection).
# CHECK      'SOURCES|STATISTIC' that must give the same number (else blank + exception), e.g. PAS-X's own
#            'Spraying Min & Max' record against the 15-minute readings.
RULE_FIELDS = ['RULE_ID', 'LAYOUT', 'SHEET', 'COLUMN', 'KIND', 'SOURCES', 'STATISTIC', 'REQUIRE', 'SCALE',
               'EXPECTED_UOM', 'MACHINE', 'ZERO_AS', 'CHECK', 'STATUS', 'EVIDENCE', 'VERSION']
RULESET_VERSION = {'F1M00011': 'F1M00011-R2 (30-Oct-2026)'}


def _rule(rid, sheet, col, kind, sources='', stat='', require='ANY', scale='', uom='', machine='', zero='',
          check='', status='VERIFIED', evidence='', layout='F1M00011'):
    return OrderedDict(RULE_ID=rid, LAYOUT=layout, SHEET=sheet, COLUMN=col, KIND=kind, SOURCES=sources,
                       STATISTIC=stat, REQUIRE=require, SCALE=s(scale), EXPECTED_UOM=uom, MACHINE=s(machine),
                       ZERO_AS=zero, CHECK=check, STATUS=status if kind in ('PASX', 'HEADER', 'FORMULA') else '',
                       EVIDENCE=evidence, VERSION=RULESET_VERSION.get(layout, layout))


def no_ave(rules, layout):
    """QA decision 30-Oct-2026: the curator writes no Ave (average) value. A rule whose QA column ends in 'Ave' is
    kept in the registry as QA_ENTRY, so the column stays in QA's layout and is left for QA."""
    leaf = {}
    for sh in layout:
        for col, p in layout_paths(sh).items():
            leaf[(sh['name'], col)] = p[-1] if p else ''
    for r in rules:
        if r['KIND'] in ('PASX', 'HEADER') and leaf.get((r['SHEET'], r['COLUMN'])) in ('ave', 'avg', 'average', 'ave.'):
            r.update(KIND='QA_ENTRY', SOURCES='', STATISTIC='', STATUS='',
                     EVIDENCE='Ave column: not filled by the curator (QA decision 30-Oct-2026). Was: ' + r['EVIDENCE'])
    return rules


def f1m00011_rules():
    return no_ave(_f1m00011_rules(), LAYOUTS['F1M00011'])


def _f1m00011_rules():
    """The approved rule set of MBR F1M00011 (30000838 Venlafaxine ER mini tablets).
    Evidence: QA's F1M00011 workbook against the real exports (66 batches Jul-Aug 2026 with v5; batch 2605012278 with
    v6.1 and its batch record report pages 42-45, 98-99)."""
    from openpyxl.utils import column_index_from_string as ci, get_column_letter as cl
    R = []
    m = 'BMR DETAILS'
    R += [_rule('BMR-SIZE', m, 'D', 'HEADER', 'BATCH_SIZE', evidence='kg from the stage name when BATCH_SIZE is a '
                'tablet count; 66/66'),
          _rule('BMR-MFG', m, 'E', 'HEADER', 'MFG_MONTH', evidence='60/66; QA uses another date in 6 batches'),
          _rule('BMR-EXP', m, 'F', 'HEADER', 'EXP_MONTH', evidence='shelf life 0 in PAS-X -> NA, as QA writes'),
          _rule('BMR-BMRNO', m, 'G', 'HEADER', 'BMR_NO', evidence='F1M00011/00000005 = QA'),
          _rule('BMR-SFG', m, 'I', 'HEADER', 'MATERIAL_CODE', evidence='66/66')]
    R += [_rule('BMR-%s' % c, m, c, 'QA_ENTRY') for c in 'HJKLMNO']

    g = 'Granulation batch size 297.52'
    R += [_rule('GRA-EQ-DISP', g, 'D', 'NO_PASX_SOURCE', evidence='dispensing equipment not found in the batch record'),
          _rule('GRA-EQ-GRAN', g, 'E', 'PASX', 'GRAN > *|@Rapid Mixer Granulator;GRAN > *|@Fluid Bed Equipment', 'TEXT',
                status='PROVISIONAL', evidence='batch record p.~20: PD/01/RMG06, PD/05/FBE03 = QA; needs SQL v6.2'),
          _rule('GRA-EQ-BLEND', g, 'F', 'PASX', 'LUB > *|@Conta Blender ID', 'TEXT', status='PROVISIONAL',
                evidence='batch record: PD/03/BLN13 = QA; needs SQL v6.2')]
    codes = {'LOT1': dict(dry='DRY1', mil='DRYMIL1', load='IMLOADE'), 'LOT2': dict(dry='DRY2', mil='DRYMIL2', load='IMPLOADE')}
    for lot, start in (('LOT1', 'G'), ('LOT2', 'AD')):
        o, k, n = ci(start), codes[lot], lot[-1]
        L = lambda i: cl(o + i)
        ipc = 'GRAN > %s > IPC' % lot
        R += [_rule('GRA-L%s-BINTIME' % n, g, L(0), 'PASX', 'GRAN > %s > WETGRAN > BINADD|MIXTM' % lot, 'VALUE',
                    uom='sec', evidence='66/66'),
              _rule('GRA-L%s-BINIMP' % n, g, L(1), 'PASX', 'GRAN > %s > WETGRAN > BINADD|@Impeller Speed' % lot, 'TEXT',
                    status='PROVISIONAL', evidence='text setting Slow in the batch record; needs SQL v6.2'),
              _rule('GRA-L%s-BINCHOP' % n, g, L(2), 'PASX', 'GRAN > %s > WETGRAN > BINADD|@Chopper Speed' % lot, 'TEXT',
                    status='PROVISIONAL', evidence='text setting; needs SQL v6.2'),
              _rule('GRA-L%s-KNTIME' % n, g, L(3), 'PASX', 'GRAN > %s > WETGRAN > MIXING|MIXTM' % lot, 'VALUE',
                    uom='sec', evidence='66/66'),
              _rule('GRA-L%s-KNIMP' % n, g, L(4), 'PASX', 'GRAN > %s > WETGRAN > MIXING|@Impeller Speed' % lot, 'TEXT',
                    status='PROVISIONAL', evidence='text setting Fast; needs SQL v6.2'),
              _rule('GRA-L%s-KNCHOP' % n, g, L(5), 'PASX', 'GRAN > %s > WETGRAN > MIXING|@Chopper Speed' % lot, 'TEXT',
                    status='PROVISIONAL', evidence='text setting Fast; needs SQL v6.2'),
              _rule('GRA-L%s-LOAD' % n, g, L(6), 'PASX', 'GRAN > %s > WETGRAN > MIXING|%s' % (lot, k['load']), 'VALUE',
                    uom='Amp', evidence='57/66 (QA differs by 0.1-0.4 A in 9)')]
        R += [_rule('GRA-L%s-ETH%d' % (n, i), g, L(i), 'NOT_IN_MBR', evidence='no ethanol step in F1M00011; QA writes NA')
              for i in range(7, 12)]
        for i, (name, pc) in enumerate((('INLET', 'INLET'), ('PROD', 'PROD'), ('EXH', 'EXH'), ('FLAP', 'FLAP'))):
            st = 'VERIFIED' if name != 'EXH' else 'PROVISIONAL'
            ev = ('2605012278 Lot-I/II: after-first = QA (inlet 36-37 / 37-38, product 18-34 / 22-35)'
                  if name != 'EXH' else 'same rule as inlet; QA leaves this red column empty for most batches')
            R += [_rule('GRA-L%s-%s-MIN' % (n, name), g, L(12 + 2 * i), 'PASX', '%s|%s' % (ipc, pc), 'MIN_AFTER_FIRST',
                        status=st, evidence=ev),
                  _rule('GRA-L%s-%s-MAX' % (n, name), g, L(13 + 2 * i), 'PASX', '%s|%s' % (ipc, pc), 'MAX_AFTER_FIRST',
                        status=st, evidence=ev)]
        R += [_rule('GRA-L%s-DRYTIME' % n, g, L(20), 'PASX', 'GRAN > %s > *|@Actual Drying Time' % lot, 'VALUE',
                    status='PROVISIONAL', evidence='formula activity: 2605012278 78 / 70 min = QA; needs SQL v6.2'),
              _rule('GRA-L%s-LOD' % n, g, L(21), 'PASX', 'GRAN > %s > %s > AFDRY|FINLOD' % (lot, k['dry']), 'VALUE',
                    evidence='66/66'),
              _rule('GRA-L%s-MILLOD' % n, g, L(22), 'PASX', 'GRAN > %s > %s > OBS|LOD' % (lot, k['mil']), 'VALUE',
                    evidence='64/66')]
    R += [_rule('GRA-REMARK1', g, 'BA', 'QA_ENTRY'), _rule('GRA-REMARK2', g, 'BC', 'QA_ENTRY'),
          _rule('GRA-BLENDLOD', g, 'BB', 'PASX', 'LUB > LOD|LOD', 'VALUE', evidence='63/66')]

    c = 'Compression'
    R += [_rule('COM-EQ', c, 'D', 'PASX', 'COMP > *|@Select Compression MC ID.', 'TEXT', status='PROVISIONAL',
                evidence='list activity: PD/08/COM08 (73 station); QA writes PD/08/COM08 (73station)')]
    for col in 'EFHIJKLMNOPQR':
        R.append(_rule('COM-IPQC-%s' % col, c, col, 'PRINTOUT', evidence='PAS-X records only "Within tolerance range"; '
                       'the numbers are on the attached printouts (scan reader)'))
    R += [_rule('COM-UNIF-AVG', c, 'G', 'PASX', 'COMP > YD|AVGWT', 'VALUE', uom='mg', evidence='65/66'),
          _rule('COM-M30-MIN', c, 'S', 'PASX', 'COMP > COMP > 30MIN|30LHS;COMP > COMP > 30MIN|30RHS', 'MIN', 'ALL',
                uom='mg', evidence='2605012278: 689.35 = QA = PAS-X Min Mass of 30 Tab. RHS'),
          _rule('COM-M30-MAX', c, 'T', 'PASX', 'COMP > COMP > 30MIN|30LHS;COMP > COMP > 30MIN|30RHS', 'MAX', 'ALL',
                uom='mg', evidence='2605012278: 715.7 = QA = PAS-X Max Mass of 30 Tab. LHS'),
          _rule('COM-M30-AVG', c, 'U', 'PASX', 'COMP > COMP > 30MIN|30LHS;COMP > COMP > 30MIN|30RHS', 'AVG', 'ALL',
                uom='mg', evidence='2605012278: 702.35 = QA (all 48 readings)'),
          _rule('COM-M30-AVGMIN', c, 'V', 'FORMULA', '=AVERAGE(S{r}/30)', evidence="QA's own formula"),
          _rule('COM-M30-AVGMAX', c, 'W', 'FORMULA', '=AVERAGE(T{r}/30)', evidence="QA's own formula"),
          _rule('COM-REMARK1', c, 'X', 'QA_ENTRY'), _rule('COM-REMARK2', c, 'AG', 'QA_ENTRY'),
          _rule('COM-REMARK3', c, 'AP', 'QA_ENTRY')]
    for stn, start in (('75', 'Y'), ('73', 'AH')):
        o = ci(start)
        L = lambda i: cl(o + i)
        sp = 'COMP > COMP > STARTUP > PARA > %sSTNS > MULTIT' % stn
        rpm = 'COMP > COMP > STARCOM > RPMCHK > %sSTNS > MULTIT' % stn
        adj = 'COMP > SUBACT > ADJUST > %sSTNS > MULTIT' % stn
        ipc = 'COMP > COMP > 60MIPC%s' % stn
        R += [_rule('COM-%s-TURRET-MIN' % stn, c, L(0), 'PASX', '%s|TURETRPM;%s|TURETRPM;%s|TURETRPM' % (sp, rpm, adj),
                    'MIN', machine=stn, uom='rpm', evidence='73: 26/26, 75: 39/40'),
              _rule('COM-%s-TURRET-MAX' % stn, c, L(1), 'PASX', '%s|TURETRPM;%s|TURETRPM;%s|TURETRPM' % (sp, rpm, adj),
                    'MAX', machine=stn, uom='rpm', evidence='73: 26/26, 75: 39/40'),
              _rule('COM-%s-FEEDER-MIN' % stn, c, L(2), 'PASX', '%s|FEEDERPM;%s|FEEDERPM' % (sp, adj), 'MIN',
                    machine=stn, uom='rpm', evidence='75: 39/40; 2605012278 (73): 12 = QA'),
              _rule('COM-%s-FEEDER-MAX' % stn, c, L(3), 'PASX', '%s|FEEDERPM;%s|FEEDERPM' % (sp, adj), 'MAX',
                    machine=stn, uom='rpm', evidence='75: 39/40; 2605012278 (73): 12 = QA')]
        if stn == '73':
            R += [_rule('COM-73-PRE-MIN', c, L(4), 'PASX', '%s|MPCFL;%s|MPCFR' % (ipc, ipc), 'MIN', 'ALL', machine=stn,
                        uom='kN', evidence='2605012278: 1.8 = QA (13 hourly checks, LHS+RHS; start-up record excluded)'),
                  _rule('COM-73-PRE-MAX', c, L(5), 'PASX', '%s|MPCFL;%s|MPCFR' % (ipc, ipc), 'MAX', 'ALL', machine=stn,
                        uom='kN', evidence='2605012278: 2.5 = QA'),
                  _rule('COM-73-MAIN-MIN', c, L(6), 'PASX', '%s|MMCFL;%s|MMCFR' % (ipc, ipc), 'MIN', 'ALL', machine=stn,
                        uom='kN', evidence='2605012278: 12.1 = QA'),
                  _rule('COM-73-MAIN-MAX', c, L(7), 'PASX', '%s|MMCFL;%s|MMCFR' % (ipc, ipc), 'MAX', 'ALL', machine=stn,
                        uom='kN', evidence='2605012278: PAS-X 14.7 (RHS); QA wrote 14.3 = LHS only -> QA to confirm')]
        else:
            R += [_rule('COM-75-PRE-MIN', c, L(4), 'NOT_IN_MBR', evidence='75-station machine has no pre-compression '
                        'force in F1M00011; QA writes NA'),
                  _rule('COM-75-PRE-MAX', c, L(5), 'NOT_IN_MBR', evidence='as above'),
                  _rule('COM-75-MAIN-MIN', c, L(6), 'PASX', '%s|MMMCF' % ipc, 'MIN', machine=stn, uom='kN',
                        status='PROVISIONAL', evidence='hourly checks only (as the 73-station rule); verify on a v6 '
                        '75-station batch'),
                  _rule('COM-75-MAIN-MAX', c, L(7), 'PASX', '%s|MMMCF' % ipc, 'MAX', machine=stn, uom='kN',
                        status='PROVISIONAL', evidence='as above')]

    t = 'Coating'
    R += [_rule('COA-EQ-DISP', t, 'D', 'NO_PASX_SOURCE', evidence='not found in the batch record'),
          _rule('COA-EQ', t, 'E', 'NO_PASX_SOURCE', evidence='coating pan ID (PD/09/COT14) not found in the batch record')]
    R += [_rule('COA-UNIF-%s' % col, t, col, 'PRINTOUT', evidence='individual weights on the printout; Ave differs '
                'from Final Average Weight in 7/66') for col in 'FGH']
    R += [_rule('COA-AVGWT-MIN', t, 'I', 'PASX', 'COAT > WTGAIN|B', 'VALUE', uom='mg', evidence='66/66'),
          _rule('COA-AVGWT-MAX', t, 'J', 'PASX', 'COAT > WTGAIN|B', 'VALUE', uom='mg', evidence='66/66'),
          _rule('COA-REMARK1', t, 'K', 'QA_ENTRY'), _rule('COA-REMARK2', t, 'AH', 'QA_ENTRY'),
          _rule('COA-SOLMIX', t, 'L', 'PASX', 'SOLPREP > COATSOL > SOLPREP > END2|MXTM2', 'VALUE', scale=60,
                uom='Hrs.', status='PROVISIONAL', evidence='1 h -> 60 min = QA in the one batch QA filled')]
    for (lo, hi), (mn, mx), ipc, name in ((('M', 'N'), ('MNINLET', 'MXINLET'), 'INLET', 'INLET'),
                                         (('O', 'P'), ('MNEXH', 'MXEXH'), 'EXH', 'EXH'),
                                         (('Q', 'R'), ('MNBED', 'MXBED'), 'BED', 'BED'),
                                         (('S', 'T'), ('MNPAN', 'MXPAN'), 'PAN', 'PAN'),
                                         (('U', 'V'), ('MNPPRPM', 'MXPPRPM'), 'PPRPM', 'PPRPM'),
                                         (('W', 'X'), ('MNCFM', 'MXCFM'), 'CFM', 'CFM'),
                                         (('Y', 'Z'), ('MNATM', 'MXATM'), 'ATM', 'ATM'),
                                         (('AA', 'AB'), ('MNSPRAY', 'MXSPRAY'), 'SPRAY', 'SPRAY')):
        chk = 'COAT > COAT > SPRAY > IPC15|%s' % ipc
        R += [_rule('COA-SPR-%s-MIN' % name, t, lo, 'PASX', 'COAT > SPRMNMX|%s' % mn, 'VALUE', check=chk + '|MIN',
                    evidence="PAS-X 'Spraying Min & Max' record; = 15-min readings in 2605012278"),
              _rule('COA-SPR-%s-MAX' % name, t, hi, 'PASX', 'COAT > SPRMNMX|%s' % mx, 'VALUE', check=chk + '|MAX',
                    evidence="PAS-X 'Spraying Min & Max' record; = 15-min readings in 2605012278")]
    R += [_rule('COA-WTGAIN', t, 'AC', 'PASX', 'COAT > WTGAIN|WTGAIN', 'VALUE', uom='%', evidence='66/66'),
          _rule('COA-DRY-INLET-MIN', t, 'AD', 'PASX', 'COAT > COAT > DRYING > START|INLET', 'MIN', evidence='2605012278: 51 = QA'),
          _rule('COA-DRY-INLET-MAX', t, 'AE', 'PASX', 'COAT > COAT > DRYING > START|INLET', 'MAX', evidence='2605012278: 51 = QA'),
          _rule('COA-DRY-BED-MIN', t, 'AF', 'PASX', 'COAT > COAT > DRYING > START|BED', 'MIN', evidence='2605012278: 41.8 = QA'),
          _rule('COA-DRY-BED-MAX', t, 'AG', 'PASX', 'COAT > COAT > DRYING > START|BED', 'MAX', evidence='2605012278: 41.8 = QA')]

    y = 'Yield'
    R += [_rule('YLD-LUB', y, 'C', 'PASX', 'LUB > YIELD|C', 'VALUE', uom='%', evidence='66/66'),
          _rule('YLD-LUB-REJ', y, 'D', 'PASX', 'LUB > YIELD|E;LUB > YIELD|F', 'SUM', zero='NA', uom='kg',
                evidence='always 0 -> NA, as QA writes'),
          _rule('YLD-COMP', y, 'E', 'PASX', 'COMP > YD|D', 'VALUE', uom='%', evidence='66/66'),
          _rule('YLD-COMP-REJ', y, 'F', 'PASX', ';'.join('COMP > YD|%s' % k for k in ('F1', 'F2', 'F3', 'F4', 'F5', 'F7')),
                'SUM', uom='kg', evidence='sum of the kg rejects: 57/66 (QA 3-decimal vs PAS-X 2-decimal parts)'),
          _rule('YLD-COAT', y, 'G', 'PASX', 'COAT > YILED|E', 'VALUE', uom='%', evidence='66/66'),
          _rule('YLD-COAT-REJ', y, 'H', 'PASX', ';'.join('COAT > YILED|%s' % k for k in ('G2', 'G3', 'G4', 'H')), 'SUM',
                uom='kg', evidence='65/66'),
          _rule('YLD-REMARK', y, 'I', 'QA_ENTRY')]

    e = 'Temp & RH'
    for (tc, rc), path, name in ((('C', 'D'), 'GRAN > LC > PR', 'GRAN'), (('E', 'F'), 'LUB > LC > PR', 'BLEND'),
                                 (('G', 'H'), 'COMP > LC > PR', 'COMP'), (('I', 'J'), 'COAT > LC > PR', 'COAT')):
        R += [_rule('ENV-%s-T' % name, e, tc, 'PASX', '%s|TEMP' % path, 'VALUE', evidence='line clearance of the stage'),
              _rule('ENV-%s-RH' % name, e, rc, 'PASX', '%s|RH' % path, 'VALUE', evidence='line clearance of the stage')]
    R += [_rule('ENV-DISP-T', e, 'K', 'PASX', 'DSPRM > LC > PR|TEMP', 'VALUE', status='PROVISIONAL',
                evidence='dispensing RM line clearance; code to confirm on a v6 export'),
          _rule('ENV-DISP-RH', e, 'L', 'PASX', 'DSPRM > LC > PR|RH', 'VALUE', status='PROVISIONAL', evidence='as above'),
          _rule('ENV-REMARK', e, 'M', 'QA_ENTRY')]
    return R


RULESETS = {'F1M00011': f1m00011_rules}
# MBR versions the rule set was proven on (66 batches Jul-Aug 2026: versions 4 and 5). Another version may have moved
# steps or codes: its batches are not filled (exception UNKNOWN_MBR_VERSION) until QA approves the rule set for it.
APPROVED_MBR_VERSIONS = {'F1M00011': ('4', '5')}
# how an export is recognised when MBR_ID is empty or different: product codes, or step codes of that MBR
LAYOUT_SIGNS = {'F1M00011': dict(materials={'30000838'}, codes=('GRAN > LOT1 > IPC', 'COMP > COMP > 30MIN',
                                                                 'COAT > SPRMNMX'))}


def write_rules(path, rules):
    with open(path, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, RULE_FIELDS)
        w.writeheader()
        for r in rules:
            w.writerow({k: r.get(k, '') for k in RULE_FIELDS})


def read_rules(path):
    with open(path, encoding='utf-8-sig') as fh:
        rules = [OrderedDict((k, s(r.get(k))) for k in RULE_FIELDS) for r in csv.DictReader(fh)]
    bad = [r['RULE_ID'] for r in rules if r['KIND'] not in ('PASX', 'HEADER', 'FORMULA', 'QA_ENTRY', 'PRINTOUT',
                                                              'NOT_IN_MBR', 'NO_PASX_SOURCE')]
    if bad:
        raise ValueError('Unknown KIND in rules: ' + ', '.join(bad))
    return rules


# ------------------------------------------------------------------------------------------------ evaluating a rule
class Result:
    def __init__(self, value=None, status='', rows=(), calc='', problem=''):
        self.value, self.status, self.rows, self.calc, self.problem = value, status, list(rows), calc, problem


def parse_sources(text):
    out = []
    for part in [p for p in s(text).split(';') if p.strip()]:
        if part.startswith('~'):                       # generic format: '~<path regex>¦<codes>'
            path, _, par = part.rpartition('¦')
            out.append((path, par.strip()))
            continue
        path, _, par = part.partition('|')
        out.append((code(path), par.strip()))
    return out


def match(b, path, par):
    """Export rows of one source: exact step code path (or 'PATH > *' = any step below) and exact parameter code
    (or '@Description' = exact activity description)."""
    if path.startswith('~'):                           # generic format: path regex + a set of parameter codes
        rx = re.compile(path[1:])
        if par.startswith('@~'):
            name = re.compile(par[2:])
            return [r for r in b if rx.fullmatch(code(r.get('STEP_CODE_PATH'))) and
                    name.search(' '.join(s(r.get('APQR_PARAMETER')).lower().split()))]
        want = set(par.split(','))
        return [r for r in b if rx.fullmatch(code(r.get('STEP_CODE_PATH'))) and s(r.get('PARAMETER_CODE')) in want]
    if path.endswith('> *'):
        pre = path[:-3].strip() + ' >'
        ok_path = lambda p: p.startswith(pre)
    else:
        ok_path = lambda p: p == path
    if par.startswith('@'):
        want = ' '.join(par[1:].split()).lower()
        ok_par = lambda r: ' '.join(s(r.get('APQR_PARAMETER')).split()).lower() == want
    else:
        ok_par = lambda r: s(r.get('PARAMETER_CODE')) == par
    return [r for r in b if ok_path(code(r.get('STEP_CODE_PATH'))) and ok_par(r)]


def _nums(rows, field):
    return [(num(r.get(field)), r) for r in rows if num(r.get(field)) is not None]


def statistic(rows, stat):
    """(value, calculation text, problem) of one statistic over the rows of one rule."""
    if stat in ('MIN', 'MAX'):
        v = _nums(rows, 'READING_' + stat)
        if not v:
            return None, '', 'no numeric reading'
        x = (min if stat == 'MIN' else max)(a for a, _ in v)
        n = sum(int(num(r.get('READING_COUNT')) or 1) for _, r in v)
        return x, '%s of %d reading(s) in %d source row(s)' % (stat.lower(), n, len(v)), ''
    if stat in ('MIN_AFTER_FIRST', 'MAX_AFTER_FIRST'):
        base = stat.split('_')[0]
        v = _nums(rows, 'READING_' + stat)
        if v:
            x = (min if base == 'MIN' else max)(a for a, _ in v)
            return x, '%s of the readings after the first (start-up) reading' % base.lower(), ''
        if rows and all((num(r.get('READING_COUNT')) or 0) == 1 for r in rows):
            return None, '', 'only one reading - the start-up reading QA leaves out - nothing to report'
        return None, '', 'no reading after the first one'
    if stat == 'AVG':
        v = [(num(r.get('READING_AVG')), num(r.get('READING_COUNT')) or 1, r) for r in rows if num(r.get('READING_AVG')) is not None]
        if not v:
            return None, '', 'no numeric reading'
        n = sum(c for _, c, _ in v)
        return sum(a * c for a, c, _ in v) / n, 'average of %d reading(s) (count-weighted over %d source row(s))' % (n, len(v)), ''
    if stat in ('IMIN', 'IMAX', 'IAVG'):
        f = {'IMIN': 'INDIVIDUAL_MIN', 'IMAX': 'INDIVIDUAL_MAX', 'IAVG': 'INDIVIDUAL_AVG'}[stat]
        v = _nums(rows, f)
        if not v:
            return None, '', 'no individual-unit values'
        if stat == 'IAVG':
            w = [(a, num(r.get('INDIVIDUAL_COUNT')) or 1) for a, r in v]
            return sum(a * c for a, c in w) / sum(c for _, c in w), 'average of individual units', ''
        return (min if stat == 'IMIN' else max)(a for a, _ in v), stat[1:].lower() + ' of individual units', ''
    if stat == 'SUM':
        v = _nums(rows, 'READING_MAX')
        if not v:
            return None, '', 'no numeric reading'
        return sum(a for a, _ in v), 'sum of %d value(s): %s' % (len(v), ' + '.join(str(clean(a)) for a, _ in v)), ''
    if stat == 'VALUE':
        v = _nums(rows, 'READING_MAX')
        if not v:
            return None, '', 'no numeric reading'
        distinct = sorted({clean(a) for a, _ in v} | {clean(num(r.get('READING_MIN'))) for _, r in v
                                                      if num(r.get('READING_MIN')) is not None})
        if len(distinct) > 1:
            return None, '', 'one value expected, PAS-X has %s' % ', '.join(str(d) for d in distinct)
        return v[0][0], 'single value', ''
    if stat == 'TEXT':
        texts = []
        for r in rows:
            for x in s(r.get('ACTUAL_TEXT')).split(' | ') if s(r.get('ACTUAL_TEXT')) else []:
                if x.strip() and x.strip() not in texts:
                    texts.append(x.strip())
            if not s(r.get('ACTUAL_TEXT')) and num(r.get('READING_MAX')) is not None:
                t = str(clean(num(r.get('READING_MAX'))))
                if t not in texts:
                    texts.append(t)
        if not texts:
            return None, '', 'no entered text'
        return '\n'.join(texts), 'entered text', ''
    raise ValueError('unknown STATISTIC ' + stat)


def header(b, field):
    if field == 'BATCH_SIZE':
        if first(b, 'BATCH_SIZE_UOM').lower() in ('kg', 'kgs'):
            return num(first(b, 'BATCH_SIZE')), 'BATCH_SIZE (kg)'
        kg = stage_kg(b)
        return kg, 'kg in the stage name (BATCH_SIZE %s %s is a count)' % (first(b, 'BATCH_SIZE'), first(b, 'BATCH_SIZE_UOM'))
    if field == 'EXP_MONTH':
        if num(first(b, 'SHELF_LIFE')) == 0:
            return 'NA', 'shelf life 0 in PAS-X material master = no expiry (QA writes NA)'
        return first(b, 'EXP_MONTH') or None, 'EXP_MONTH'
    if field == 'BMR_NO':
        mbr, ver = first(b, 'MBR_ID'), re.sub(r'\.0+$', '', first(b, 'MBR_VERSION'))
        return (mbr + ('/' + ver.zfill(8) if ver.isdigit() else '')) or None, 'MBR_ID / MBR_VERSION'
    if field == 'MATERIAL_CODE':
        v = first(b, 'MATERIAL_CODE')
        return (int(v) if v.isdigit() else v) or None, 'MATERIAL_CODE'
    return first(b, field) or None, field


def station(b):
    """Compression machine of the batch from the MBR branch it ran ('73STNS' / '75STNS')."""
    codes = ' '.join(s(r.get('STEP_CODE_PATH')) for r in b)
    has73, has75 = bool(re.search(r'73STNS|MIPC73', codes)), bool(re.search(r'75STNS|MIPC75', codes))
    return '73' if has73 and not has75 else '75' if has75 and not has73 else ('both' if has73 else None)


REPEATED_STATS = ('MIN', 'MAX', 'AVG', 'MIN_AFTER_FIRST', 'MAX_AFTER_FIRST', 'IMIN', 'IMAX', 'IAVG')


def evaluate(rule, b):
    """Result of one rule for one batch (b = export rows of the batch)."""
    kind = rule['KIND']
    if kind == 'HEADER' and rule.get('STATUS') in ('AUTO', 'LEARNED'):
        v, how = header(b, rule['SOURCES'])
        return Result(v, 'FILLED_' + rule['STATUS'] if v not in (None, '') else 'NO_SOURCE_IN_EXPORT', calc=how)
    if kind in ('QA_ENTRY', 'PRINTOUT', 'NOT_IN_MBR', 'NO_PASX_SOURCE'):
        return Result(None, kind, calc=rule.get('EVIDENCE', ''))
    if kind == 'HEADER':
        v, how = header(b, rule['SOURCES'])
        return Result(v, 'FILLED' if v not in (None, '') else 'NO_SOURCE_IN_EXPORT', calc=how,
                      problem='' if v not in (None, '') else 'batch header field empty in the export')
    if kind == 'FORMULA':
        return Result(rule['SOURCES'], 'FORMULA', calc="QA's own formula")
    if rule.get('MACHINE'):
        st = station(b)
        if st is None:
            return Result(None, 'NO_SOURCE_IN_EXPORT', calc='machine branch not in the export')
        if st == 'both':
            return Result(None, 'EXCEPTION', problem='both 73- and 75-station steps in the export: QA to decide')
        if st != rule['MACHINE']:
            return Result('NA', 'NA_OTHER_MACHINE', calc='batch made on the %s-station machine' % st)
    srcs = parse_sources(rule['SOURCES'])
    found = [(p, q, match(b, p, q)) for p, q in srcs]
    rows = [r for _, _, rs in found for r in rs]
    missing = ['%s|%s' % (p, q) for p, q, rs in found if not rs]
    if not rows:
        return Result(None, 'NO_SOURCE_IN_EXPORT', problem='no export row for ' + ', '.join(missing))
    if rule.get('REQUIRE') == 'ALL' and missing:
        return Result(None, 'EXCEPTION', rows, problem='incomplete: no export row for ' + ', '.join(missing))
    uom = rule.get('EXPECTED_UOM')
    if uom:
        wrong = sorted({s(r.get('UOM')) for r in rows if s(r.get('UOM')).lower() != uom.lower()})
        if wrong:
            return Result(None, 'EXCEPTION', rows, problem='unit %s, rule expects %s' % (', '.join(wrong), uom))
    if rule['STATISTIC'] in REPEATED_STATS and 'READING_VALUES' not in rows[0]:
        return Result(None, 'EXCEPTION', rows, problem='export made before SQL v6: the repetitions of the step are '
                      'missing, a Min / Max / Avg from it would be wrong - run APQR_FINAL_QUERY.sql v6.2')
    corr = [r for r in rows if (num(r.get('CORRECTED_ENTRIES')) or 0) > 0]
    v, calc, problem = statistic(rows, rule['STATISTIC'])
    if v is None:
        return Result(None, 'EXCEPTION', rows, problem=problem)
    if rule.get('SCALE') and isinstance(v, float):
        v, calc = v * float(rule['SCALE']), calc + ' x %s' % rule['SCALE']
    if rule.get('CHECK'):
        cpath, cpar, cstat = (s(rule['CHECK']).split('|') + ['', '', ''])[:3]
        crows = match(b, code(cpath), cpar)
        if crows and all((num(r.get('READING_COUNT')) or 0) > 1 for r in crows):
            cv, _, _ = statistic(crows, cstat)
            if cv is None or abs(cv - v) > 1e-6 + abs(v) * 1e-9:
                return Result(None, 'EXCEPTION', rows + crows, problem='%s = %s but the readings of %s give %s %s' % (
                    rule['SOURCES'], clean(v), cpath + '|' + cpar, cstat.lower(), clean(cv)))
            calc += '; checked: %s of %s = same' % (cstat.lower(), cpath + '|' + cpar)
        elif crows:
            calc += '; check skipped (one reading only: export before v6)'
    if isinstance(v, float) and v == 0 and rule.get('ZERO_AS'):
        v, calc = rule['ZERO_AS'], calc + '; 0 -> %s' % rule['ZERO_AS']
    status = {'VERIFIED': 'FILLED', 'AUTO': 'FILLED_AUTO', 'LEARNED': 'FILLED_LEARNED'}.get(rule.get('STATUS') or
                                                                                           'VERIFIED', 'FILLED_PROVISIONAL')
    res = Result(clean(v), status, rows, calc)
    if corr:
        res.problem = 'source has corrected entries (%s): review' % '; '.join(s(r.get('SUPERSEDED_VALUES')) for r in corr)
    if missing:
        res.calc += ' (sources without rows: %s)' % ', '.join(missing)
    return res


# ------------------------------------------------------------------------------------------------ workbook
def _style_header(cell, fill):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    side = Side(style='thin')
    cell.font = Font(name='Times New Roman', size=10, bold=True)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = Border(left=side, right=side, top=side, bottom=side)
    if fill:
        cell.fill = PatternFill('solid', fgColor='FF' + fill)


def _style_data(cell):
    from openpyxl.styles import Alignment, Border, Font, Side
    side = Side(style='thin')
    cell.font = Font(name='Times New Roman', size=10)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = Border(left=side, right=side, top=side, bottom=side)


def builtin_workbook(layout):
    """QA's header rows (built-in copy) as a workbook: what the template looks like without its sample rows."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in layout:
        ws = wb.create_sheet(sh['name'][:31])
        for r in range(1, sh['first']):
            for c in range(1, sh['ncols'] + 1):
                _style_header(ws.cell(r, c), None)
        for coord, v, fill in sh['cells']:
            ws[coord].value = v
            _style_header(ws[coord], fill)
        for m in sh['merges']:
            ws.merge_cells(m)
        for col, w in sh['widths'].items():
            ws.column_dimensions[col].width = w
        for r, h in sh['heights'].items():
            ws.row_dimensions[int(r)].height = h
        for c in range(1, sh['ncols'] + 1):
            if get_column_letter(c) not in sh['widths']:
                ws.column_dimensions[get_column_letter(c)].width = 11
        ws.freeze_panes = ws.cell(sh['first'], (sh['bcol'] or 1) + 1)
    return wb


def _norm(v):
    return ' '.join(s(v).replace('ˉ', '¯').split()).lower()


def layout_paths(sh):
    from openpyxl.utils.cell import coordinate_from_string
    paths = {}
    for coord, v, _fill in sh['cells']:
        col, row = coordinate_from_string(coord)
        if v is not None and s(v):
            paths.setdefault(col, []).append((row, _norm(v)))
    return {c: tuple(x for _, x in sorted(p)) for c, p in paths.items()}


def template_paths(ws, first_row):
    from openpyxl.utils import get_column_letter
    paths = {}
    for (r, c), cell in ws._cells.items():                  # QA sheets are 16,384 columns wide: only stored cells
        if r < first_row and type(cell).__name__ != 'MergedCell' and cell.value is not None and s(cell.value):
            paths.setdefault(get_column_letter(c), []).append((r, _norm(cell.value)))
    return {c: tuple(x for _, x in sorted(p)) for c, p in paths.items()}


def header_check(sh, ws):
    """(identical, differing columns) of the template's header rows against the rule set's layout. The rule set
    names columns by letter, so a template whose headers moved is refused rather than filled into wrong columns."""
    lp, tp = layout_paths(sh), template_paths(ws, sh['first'])
    diff = sorted((c for c, p in lp.items() if tp.get(c) != p), key=lambda c: (len(c), c))
    return len(lp) - len(diff), diff


REF = re.compile(r"^=\s*\+?\(?\s*'?([^'!()]+)'?!\$?([A-Z]{1,3})\$?(\d+)\s*\)?\s*$")   # ='X'!C5, =('X'!C5)


def resolved(wb, value, depth=0):
    """Value of a cell, following QA's links ='BMR DETAILS'!C2055 (other sheets repeat the master's batch number)."""
    m = REF.match(value) if isinstance(value, str) else None
    if not m or depth > 3:
        return value
    if m.group(1) not in wb.sheetnames:
        return None
    return resolved(wb, wb[m.group(1)]['%s%s' % (m.group(2), m.group(3))].value, depth + 1)


def batch_rows(wb, ws, col, first_row):
    """{batch number: row} of a sheet (links resolved), the last row holding a batch, and that row's batch cell."""
    from openpyxl.utils import column_index_from_string as ci
    c = ci(col)
    rows, last, last_cell = {}, first_row - 1, None
    for (r, cc), cell in list(ws._cells.items()):
        if cc != c or r < first_row:
            continue
        v = batch_no(resolved(wb, cell.value))
        if v:
            rows.setdefault(v, r)
            if r > last:
                last, last_cell = r, cell.value
    return rows, last, last_cell


def link_row(wb, ws, col, first_row, batch):
    from openpyxl.utils import column_index_from_string as ci
    c = ci(col)
    for (r, cc), cell in list(ws._cells.items()):
        if cc == c and r >= first_row and REF.match(s(cell.value)) and batch_no(resolved(wb, cell.value)) == batch:
            return r
    return None


def field_name(sh, col):
    p = layout_paths(sh).get(col)
    return ' > '.join(p) if p else col


def fill_sheet(wb, ws, sh, rules, batches, order, trace, exc, master_rows, from_template):
    """Write the batches of one sheet: the master sheet (BMR DETAILS) gets the batch number; the other sheets link
    to the master row the way QA's workbook does (='BMR DETAILS'!C7) when the template does so."""
    from openpyxl.utils import column_index_from_string as ci, get_column_letter
    bcol = get_column_letter(sh['bcol'])
    have, last, last_cell = batch_rows(wb, ws, bcol, sh['first'])
    link = REF.match(s(last_cell)) if last_cell is not None else None
    style_row = last if last >= sh['first'] else None
    srcol = 1 if sh['bcol'] > 1 else None
    for b in order:
        new = b not in have
        r = have.get(b) or link_row(wb, ws, bcol, sh['first'], b)
        if r is None:
            last += 1
            r = last
            for c in range(1, sh['ncols'] + 1):
                if style_row and ws.cell(style_row, c).has_style:
                    ws.cell(r, c)._style = copy.copy(ws.cell(style_row, c)._style)
                elif not style_row:
                    _style_data(ws.cell(r, c))
            if srcol and ws.cell(r, srcol).value in (None, ''):                 # Sr. No. continues the numbering
                prev = ws.cell(r - 1, srcol).value if r - 1 >= sh['first'] else 0
                ws.cell(r, srcol).value = (int(prev) if isinstance(prev, (int, float)) else r - sh['first']) + 1
            if link and b in master_rows:                                       # QA's own link to the master row
                f = link.group(0)
                ws.cell(r, sh['bcol']).value = f[:link.start(3)] + str(master_rows[b]) + f[link.end(3):]
            else:
                ws.cell(r, sh['bcol']).value = int(b) if b.isdigit() else b
        have[b] = r
        for rule in rules:
            cell = ws['%s%d' % (rule['COLUMN'], r)]
            res = evaluate(rule, batches[b])
            if res.status == 'FORMULA':
                res.value = rule['SOURCES'].replace('{r}', str(r))
            if cell.value not in (None, '') and res.status not in ('QA_ENTRY', 'PRINTOUT', 'NOT_IN_MBR', 'NO_PASX_SOURCE'):
                if not (new and from_template is False):
                    trace.append(trace_row(b, batches[b], sh, rule, cell.coordinate, cell.value,
                                           Result(cell.value, 'KEPT', calc='template cell already filled: not overwritten')))
                    continue
            if res.value is not None and res.status in ('FILLED', 'FILLED_PROVISIONAL', 'FILLED_AUTO', 'FILLED_LEARNED', 'FORMULA', 'NA_OTHER_MACHINE'):
                cell.value = res.value
            trace.append(trace_row(b, batches[b], sh, rule, cell.coordinate, cell.value if res.value is not None else None, res))
            if res.status == 'EXCEPTION' or (res.problem and res.status in ('FILLED', 'FILLED_PROVISIONAL', 'FILLED_AUTO', 'FILLED_LEARNED')) or \
                    (res.status == 'NO_SOURCE_IN_EXPORT' and rule.get('STATUS') == 'VERIFIED' and rule['KIND'] == 'PASX'):
                exc.append([b, sh['name'], cell.coordinate, field_name(sh, rule['COLUMN']), rule['RULE_ID'],
                            res.status, res.problem or 'expected source not in the export'])
    return have


TRACE_COLS = ['BATCH_NO', 'MATERIAL_CODE', 'MBR_ID', 'MBR_VERSION', 'APQR_SHEET', 'OUTPUT_CELL', 'APQR_FIELD',
              'OUTPUT_VALUE', 'STATUS', 'RULE_ID', 'RULE_STATUS', 'SOURCE_STEP_CODE_PATH', 'SOURCE_PARAMETER_CODE',
              'SOURCE_PARAMETER', 'SOURCE_STATISTIC', 'SOURCE_READING_COUNT', 'SOURCE_READING_VALUES',
              'SOURCE_TEXT', 'SOURCE_FIRST_READING', 'SOURCE_LAST_READING', 'CALCULATION', 'PROBLEM', 'MAPPING_VERSION',
              'SOURCE_TYPE']
# SOURCE_TYPE: PASX_DIRECT (one value as entered), PASX_DERIVED (min / max / avg / sum over readings), PASX_BATCH_HEADER,
# QA_TEMPLATE_FORMULA (QA's own formula), RULE_NA (approved rule: other machine / no rejection); never a free constant


def trace_row(b, rows, sh, rule, cell, value, res):
    j = lambda f: ' ;; '.join(dict.fromkeys(s(r.get(f)) for r in res.rows if s(r.get(f))))
    return [first(rows, 'BATCH_NO'), first(rows, 'MATERIAL_CODE'), first(rows, 'MBR_ID'), first(rows, 'MBR_VERSION'),
            sh['name'], cell, field_name(sh, rule['COLUMN']), value, res.status, rule['RULE_ID'], rule.get('STATUS', ''),
            j('STEP_CODE_PATH'), j('PARAMETER_CODE'), j('APQR_PARAMETER'), rule.get('STATISTIC', ''),
            sum(int(num(r.get('READING_COUNT')) or 0) for r in res.rows) if res.rows else None,
            j('READING_VALUES')[:1000], j('ACTUAL_TEXT')[:500], j('FIRST_READING_ON'), j('LAST_READING_ON'),
            res.calc, res.problem, rule.get('VERSION', ''), source_type(rule, res, value)]


def source_type(rule, res, value):
    if value in (None, '') or res.status not in ('FILLED', 'FILLED_PROVISIONAL', 'FILLED_AUTO', 'FILLED_LEARNED',
                                                 'FORMULA', 'NA_OTHER_MACHINE'):
        return ''
    if res.status == 'FORMULA':
        return 'QA_TEMPLATE_FORMULA'
    if res.status == 'NA_OTHER_MACHINE' or (rule.get('ZERO_AS') and value == rule.get('ZERO_AS')):
        return 'RULE_NA'
    if rule['KIND'] == 'HEADER':
        return 'PASX_BATCH_HEADER'
    if not res.rows:
        return ''
    return 'PASX_DIRECT' if rule.get('STATISTIC') in ('VALUE', 'TEXT') else 'PASX_DERIVED'


def _table(ws, head, rows, widths):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws.append(head)
    for c in range(1, len(head) + 1):
        ws.cell(1, c).font = Font(bold=True, color='FFFFFF')
        ws.cell(1, c).fill = PatternFill('solid', fgColor='FF0B4F8A')
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def own_sheets(wb, trace, exc, rules, info, validation=None, learned=None, unmapped_rows=None):
    for n in OWN_SHEETS:
        if n in wb.sheetnames:
            del wb[n]
    ws = wb.create_sheet(TRACE_SHEET)
    _table(ws, TRACE_COLS, trace, [12, 10, 10, 6, 22, 8, 40, 12, 18, 20, 12, 34, 14, 26, 16, 8, 40, 24, 16, 16, 40, 40, 22, 20])
    ws = wb.create_sheet(EXC_SHEET)
    _table(ws, ['BATCH_NO', 'APQR_SHEET', 'OUTPUT_CELL', 'APQR_FIELD', 'RULE_ID', 'STATUS', 'REASON'], exc,
           [12, 24, 10, 50, 22, 18, 90])
    if validation is not None:
        ws = wb.create_sheet(VALID_SHEET)
        _table(ws, ['BATCH_NO', 'APQR_SHEET', 'CELL', 'APQR_FIELD', 'QA_VALUE', 'CURATED_VALUE', 'RESULT', 'RULE_ID',
                    'LIKELY_CAUSE'], validation, [12, 24, 8, 50, 14, 14, 12, 22, 90])
    if learned is not None:
        ws = wb.create_sheet(LEARN_SHEET)
        _table(ws, ['QA_SHEET', 'COLUMN', 'QA_HEADER', 'RESULT', 'PASX_SOURCE_STATISTIC', 'AGREED', 'OF_QA_VALUES'],
               learned, [24, 8, 60, 14, 60, 8, 8])
    if unmapped_rows:
        ws = wb.create_sheet(UNMAPPED_SHEET)
        _table(ws, ['APQR_SHEET', 'STAGE', 'STEP_PATH', 'PARAMETER', 'UOM', 'STEP_CODE_PATH', 'PARAMETER_CODE',
                    'BATCHES_WITH_DATA', 'EXAMPLE'], unmapped_rows, [18, 22, 40, 30, 8, 34, 14, 8, 30])
    ws = wb.create_sheet(RULE_SHEET)
    for k, v in info:
        ws.append([k, v])
    ws.append([])
    start = ws.max_row + 1
    ws.append(RULE_FIELDS)
    for r in rules:
        ws.append([r.get(k, '') for k in RULE_FIELDS])
    from openpyxl.styles import Font
    for c in range(1, len(RULE_FIELDS) + 1):
        ws.cell(start, c).font = Font(bold=True)
    for i, w in enumerate([22, 10, 24, 8, 14, 60, 16, 8, 6, 8, 8, 6, 40, 12, 70, 26], 1):
        ws.column_dimensions[chr(64 + i)].width = w


# ------------------------------------------------------------------------------------------------ any product
# A product without an approved rule set is no longer refused. Two ways, both from exact PAS-X codes only:
#   GENERIC  the generic QA format (see GENERIC below): QA's sheets and header levels, the same columns for every
#            product, filled only where PAS-X has the field's exact code pattern. (Formerly: one column group per
#            STEP_CODE_PATH|PARAMETER_CODE, QA's header style (step > parameter > limit > Min / Max / Ave).
#   LEARNED  QA's own workbook for the product is given (--learn): every QA column gets the one PAS-X series and
#            statistic that reproduces QA's values in the batches both files have (>= 80 %, tolerance 0.1 %);
#            the output is QA's exact header block with the requested batches. Where several sources would write
#            different values the column stays blank (AMBIGUOUS); where none reproduces QA, NOT_LEARNED.
# Values are never estimated: only where they are written differs. Both are marked in every trace row.
SHEET_ORDER = ('dispens', 'sift', 'granul', 'blend', 'lubric', 'compress', 'capsul', 'coat', 'yield', 'temp', 'ip',
               'bpr', 'pack', 'fp')
STAT_LABEL = OrderedDict([('VALUE', 'Value'), ('MIN', 'Min'), ('MAX', 'Max'), ('AVG', 'Ave'), ('TEXT', 'Entry'),
                          ('IMIN', 'Unit Min'), ('IMAX', 'Unit Max'), ('IAVG', 'Unit Ave')])
HEAD_FILL = ('8EB4E3', 'DDD9C3', 'DBEEF4', 'F2F2F2', 'FFFFFF')
MASTER_FIELDS = [('Batch Size (In Kg)', 'BATCH_SIZE'), ('Mfg. Date', 'MFG_MONTH'), ('Exp. Date', 'EXP_MONTH'),
                 ('BMR No.', 'BMR_NO'), ('Material code', 'MATERIAL_CODE'), ('Product', 'PRODUCT_NAME'),
                 ('Mfg. order', 'MANUFACTURING_ORDER'), ('Mfg. start', 'MFG_START')]


def series(batches):
    """Every PAS-X series of the export: {(code path, parameter code or @description): facts}, in step order."""
    out = OrderedDict()
    for rows in sorted(batches.values(), key=len, reverse=True):          # the fullest batch sets the order
        for r in rows:
            path = code(r.get('STEP_CODE_PATH'))
            par = s(r.get('PARAMETER_CODE')) or '@' + ' '.join(s(r.get('APQR_PARAMETER')).split())
            if not path or par == '@':
                continue
            f = out.setdefault((path, par), dict(sheet=s(r.get('APQR_SHEET')) or s(r.get('STAGE')) or 'Other',
                                                 stage=s(r.get('STAGE')), step=s(r.get('STEP_PATH')) or s(r.get('STEP')),
                                                 param=s(r.get('APQR_PARAMETER')) or par.lstrip('@'), uom=s(r.get('UOM')),
                                                 lo=s(r.get('SPEC_MIN')), hi=s(r.get('SPEC_MAX')),
                                                 target=s(r.get('TARGET_SET_VALUE')), numeric=False, text=False,
                                                 repeated=False, units=False))
            f['numeric'] |= num(r.get('READING_MAX')) is not None
            f['text'] |= bool(s(r.get('ACTUAL_TEXT')))
            f['repeated'] |= (num(r.get('READING_COUNT')) or 0) > 1
            f['units'] |= (num(r.get('INDIVIDUAL_COUNT')) or 0) > 0
            for k, fld in (('lo', 'SPEC_MIN'), ('hi', 'SPEC_MAX'), ('uom', 'UOM')):
                f[k] = f[k] or s(r.get(fld))
    return out


def series_stats(f):
    st = (['MIN', 'MAX', 'AVG'] if f['repeated'] else ['VALUE']) if f['numeric'] else ['TEXT'] if f['text'] else []
    return st + (['IMIN', 'IMAX', 'IAVG'] if f['units'] else [])


def _source(path, par):
    return '%s|%s' % (path, par)


def _limit_text(f):
    lo, hi = clean(num(f['lo'])) if num(f['lo']) is not None else f['lo'], clean(num(f['hi'])) if num(f['hi']) is not None else f['hi']
    t = '%s - %s' % (lo, hi) if lo != '' and hi != '' else 'NLT %s' % lo if lo != '' else 'NMT %s' % hi if hi != '' else ''
    if f['target'] and not t:
        t = 'Set %s' % f['target']
    return (t + (' ' + f['uom'] if t and f['uom'] else '')).strip() or (f['uom'] or '-')


def _sheet_rank(name):
    low = name.lower()
    return next((i for i, k in enumerate(SHEET_ORDER) if k in low), len(SHEET_ORDER))


# ------------------------------------------------------------------------------------------------ generic QA format
# One APQR format in QA's style for every product without an approved rule set: QA's sheets, QA's header levels
# (group > part > parameter > limit > Min / Max), the same columns for every product. A group repeats for each lot,
# coating layer or machine the export has. A cell is filled only when PAS-X has the exact code pattern of its field
# (stage root + step segment + parameter code); otherwise it stays blank. No Ave columns (QA decision 30-Oct-2026).
# Patterns are plant code conventions seen in F1M00011 and F1M00207: GRAN > LOTn | GRANn, ... > IPC|INLET, COMP >
# 60MIPC73 | FORCE61 |MMCFL | MCFLHS, COATxxx > ... SPRAY... |INLET, <stage> > LC > PR|TEMP. QA reviews the format once.
LOT_RE = r'(?:LOT|GRAN)\d+'
SPRAY_IPC = r'(?:.+ > )?SPRAY\w* > (?:IPC\w*|SPRAYIPC)'
GENERIC = [
    dict(sheet='Granulation', roots=r'GRAN', lots=True, fields=[
        ('Dry mixing', 'Mixing time', r'(?:.+ > )?DRYMIX(?: > .+)?', 'MIXTM', ('VALUE',)),
        ('Binder addition', 'Mixing time', r'(?:.+ > )?BINADD', 'MIXTM', ('VALUE',)),
        ('Wet mixing', 'Kneading time', r'(?:.+ > )?MIXING', 'MIXTM', ('VALUE',)),
        ('Wet mixing', 'Impeller load at end', r'(?:.+ > )?MIXING', 'IMLOADE,IMPLOADE,IMPLOAD', ('VALUE',)),
        ('Drying', 'Inlet temperature', r'(?:DRY\w* > )?IPC\w*', 'INLET', ('MIN', 'MAX')),
        ('Drying', 'Product temperature', r'(?:DRY\w* > )?IPC\w*', 'PROD,PRODUCT,PRD', ('MIN', 'MAX')),
        ('Drying', 'Exhaust temperature', r'(?:DRY\w* > )?IPC\w*', 'EXH,EXHAUST', ('MIN', 'MAX')),
        ('Drying', '% LOD after drying', r'(?:DRY\w* > )?AFDRY', 'FINLOD', ('VALUE',)),
        ('Milling', '% LOD of milled granules', r'(?:.+ > )?(?:DRYMIL|MIL)\w*(?: > .+)?', 'LOD', ('VALUE',))]),
    dict(sheet='Granulation', roots=r'LUB|BLEN|BLEND', lots=False, fields=[
        ('Blending', 'Blending time', r'(?:.+ > )?BLEND > END', 'END,MIXTM', ('VALUE',)),
        ('Blending', 'Blender RPM', r'(?:.+ > )?BLEND > START', 'RPM', ('VALUE',)),
        ('Blending', '% LOD', r'LOD', 'LOD', ('VALUE',))]),
    dict(sheet='Compression', roots=r'COMP', lots=False, fields=[
        ('Machine set-up', 'Turret speed', r'.+', 'TURETRPM', ('MIN', 'MAX')),
        ('Machine set-up', 'Feeder speed', r'.+', 'FEEDERPM', ('MIN', 'MAX')),
        ('In-process checks', 'Pre-compression force', r'COMP > (?:\d+MIPC\w*|FORCE\w*)', 'MPCFL,MPCFR,PCFLHS,PCFRHS',
         ('MIN', 'MAX')),
        ('In-process checks', 'Main compression force', r'COMP > (?:\d+MIPC\w*|FORCE\w*)',
         'MMCFL,MMCFR,MMMCF,MCFLHS,MCFRHS', ('MIN', 'MAX')),
        ('In-process checks', 'Mass of tablets (LHS + RHS)', r'COMP > \d+MIN\w*',
         '10LHS,10RHS,20LHS,20RHS,30LHS,30RHS', ('MIN', 'MAX')),
        ('Lubricated blend', '% LOD', r'LOD', 'LOD', ('VALUE',))]),
    dict(sheet='Coating', roots=r'COAT\w*', lots=False, fields=[
        ('Pre-heating', 'Inlet air temperature', r'(?:.+ > )?PREHEAT(?: > .+)?', 'INLET', ('MIN', 'MAX')),
        ('Pre-heating', 'Bed temperature', r'(?:.+ > )?PREHEAT(?: > .+)?', 'BED', ('MIN', 'MAX')),
        ('Spraying', 'Inlet air temperature', SPRAY_IPC, 'INLET', ('MIN', 'MAX')),
        ('Spraying', 'Exhaust air temperature', SPRAY_IPC, 'EXH', ('MIN', 'MAX')),
        ('Spraying', 'Bed temperature', SPRAY_IPC, 'BED', ('MIN', 'MAX')),
        ('Spraying', 'Atomising air pressure', SPRAY_IPC, 'ATM', ('MIN', 'MAX')),
        ('Spraying', 'Pan RPM', SPRAY_IPC, 'PAN', ('MIN', 'MAX')),
        ('Spraying', 'Peristaltic pump RPM', SPRAY_IPC, 'PPRPM', ('MIN', 'MAX')),
        ('Spraying', 'Spray rate', SPRAY_IPC, 'SPRAY,1SPRAY', ('MIN', 'MAX')),
        ('Drying', 'Inlet air temperature', r'(?:.+ > )?DRYING(?: > .+)?', 'INLET', ('MIN', 'MAX')),
        ('Drying', 'Bed temperature', r'(?:.+ > )?DRYING(?: > .+)?', 'BED', ('MIN', 'MAX')),
        ('Cooling', 'Inlet air temperature', r'(?:.+ > )?COOLING(?: > .+)?', 'INLET', ('MIN', 'MAX')),
        ('Weight gain', '% weight gain', r'(?:.+ > )?WTGAIN', 'WTGAIN', ('VALUE',))]),
    dict(sheet='Yield', roots=r'\w+', lots=False, only_found=True, fields=[
        ('Yield', '% Actual yield', r'(?:YD|YIELD|YILED|SOLRE)', '@~^% actual yield', ('VALUE',))]),
    dict(sheet='Temp & RH', roots=r'\w+', lots=False, only_found=True, fields=[
        ('Line clearance', 'Temperature', r'LC > PR', 'TEMP', ('VALUE',)),
        ('Line clearance', '% RH', r'LC > PR', 'RH', ('VALUE',))])]
GENERIC_STAT = {'VALUE': 'Value', 'MIN': 'Min', 'MAX': 'Max'}
GENERIC_SHEETS = ('BMR DETAILS', 'Granulation', 'Compression', 'Coating', 'Yield', 'Temp & RH')


def _gsource(prefix, rest, codes):
    return '~%s > (?:%s)¦%s' % (re.escape(prefix), rest, codes)


def generic_layout(batches, mbr='GENERIC'):
    """(layout, rules) of the generic QA format for any product (see GENERIC)."""
    from openpyxl.utils import get_column_letter as cl
    lname = 'GENERIC-' + (mbr or 'unknown')
    RULESET_VERSION.setdefault(lname, "generic QA format (curator %s): QA's layout, filled where PAS-X has the field's "
                                      "exact code pattern - QA to review" % VERSION)
    ser = series(batches)
    rules, lay = [], []
    cells, merges = [('A1', 'BMR details - generic QA format (%s), not yet approved by QA' % mbr, HEAD_FILL[0]),
                      ('A2', 'Sr. No.', HEAD_FILL[1]), ('B2', 'Batch No.', HEAD_FILL[1])], ['A2:A6', 'B2:B6']
    master = MASTER_FIELDS[:5] + [('Ref. Doc No.', None), ('SFG / material code', 'MATERIAL_CODE'),
                                  ('Product', 'PRODUCT_NAME'), ('Market', None), ('Remarks', None)]
    master = [m for m in master if m[0] != 'Material code']
    widths = {'A': 7, 'B': 13}
    for i, (label, field) in enumerate(master):
        c = cl(3 + i)
        cells.append(('%s2' % c, label, HEAD_FILL[1]))
        merges.append('%s2:%s6' % (c, c))
        widths[c] = 15
        rules.append(_rule('GEN-BMR-%s' % c, 'BMR DETAILS', c, 'HEADER' if field else 'QA_ENTRY', field or '',
                           status='AUTO', layout=lname, evidence='batch header field of the export' if field else
                           'QA enters it'))
    n = 2 + len(master)
    merges.append('A1:%s1' % cl(n))
    lay.append(dict(name='BMR DETAILS', first=7, bcol=2, ncols=n, cells=cells, merges=merges, widths=widths,
                    heights={1: 20}))
    by_sheet = OrderedDict((n_, []) for n_ in GENERIC_SHEETS[1:])
    for g in GENERIC:
        inst = OrderedDict()
        for (path, par), f in ser.items():
            seg = path.split(' > ')
            if not re.fullmatch(g['roots'], seg[0]):
                continue
            pre = seg[0] + (' > ' + seg[1] if g['lots'] and len(seg) > 2 and re.fullmatch(LOT_RE, seg[1]) else '')
            rest = path[len(pre) + 3:]
            for fld in g['fields']:
                ok_par = (re.search(fld[3][2:], ' '.join(f['param'].lower().split())) if fld[3].startswith('@~')
                          else par in fld[3].split(','))
                if ok_par and re.fullmatch(fld[2], rest):
                    inst.setdefault(pre, dict(stage=f['stage'], hits={}))['hits'].setdefault(fld[:2], f)
        if not inst and not g.get('only_found'):
            inst[g['roots'].split('|')[0]] = dict(stage=g['sheet'], hits={})
        for pre, i in sorted(inst.items(), key=lambda kv: kv[0]):
            by_sheet[g['sheet']].append((g, pre, i))
    for sheet, groups in by_sheet.items():
        cells = [('A2', 'Sr. No.', HEAD_FILL[1]), ('B2', 'Batch No.', HEAD_FILL[1])]
        merges, widths, c = ['A2:A6', 'B2:B6'], {'A': 7, 'B': 13}, 3
        for g, pre, i in groups:
            lot = pre.split(' > ')[1] if ' > ' in pre else ''
            label = (i['stage'] or pre) + (' - ' + lot if lot else '') + ('  [%s]' % pre.split(' > ')[0]
                                                                       if g['roots'] != r'GRAN' else '')
            g0 = c
            part_start = {}
            for part, param, rest, codes, stats in g['fields']:
                f = i['hits'].get((part, param))
                a, z = cl(c), cl(c + len(stats) - 1)
                if part not in part_start:
                    part_start[part] = c
                    cells.append(('%s3' % a, part, HEAD_FILL[2]))
                cells += [('%s4' % a, param, HEAD_FILL[2]), ('%s5' % a, _limit_text(f) if f else '-', HEAD_FILL[3])]
                if len(stats) > 1:
                    merges += ['%s4:%s4' % (a, z), '%s5:%s5' % (a, z)]
                for k, st in enumerate(stats):
                    col = cl(c + k)
                    if len(stats) > 1:                       # Min / Max; a single value has no sub-header (as QA)
                        cells.append(('%s6' % col, GENERIC_STAT[st], HEAD_FILL[4]))
                    else:
                        merges.append('%s5:%s6' % (col, col))
                    widths[col] = 11
                    rules.append(_rule('GEN-%s-%s' % (re.sub(r'\W+', '', sheet)[:6], col), sheet, col, 'PASX',
                                       _gsource(pre, rest, codes), st, status='AUTO', layout=lname,
                                       evidence='generic field %s > %s: code pattern %s | %s' % (part, param, rest, codes)))
                c += len(stats)
            parts = list(part_start.items())
            for k, (part, start) in enumerate(parts):
                end = (parts[k + 1][1] if k + 1 < len(parts) else c) - 1
                if end > start:
                    merges.append('%s3:%s3' % (cl(start), cl(end)))
            cells.append(('%s2' % cl(g0), label, HEAD_FILL[1]))
            if c - 1 > g0:
                merges.append('%s2:%s2' % (cl(g0), cl(c - 1)))
        cells.insert(0, ('A1', '%s - generic QA format, not yet approved by QA' % sheet, HEAD_FILL[0]))
        merges.append('A1:%s1' % cl(max(c - 1, 2)))
        lay.append(dict(name=sheet, first=7, bcol=2, ncols=max(c - 1, 2), cells=cells, merges=merges, widths=widths,
                        heights={2: 40, 3: 30, 4: 40, 5: 30}))
    return lay, rules


# ------------------------------------------------------------------------------------------------ learning
LEARN_STATS = ('VALUE', 'MIN', 'MAX', 'AVG', 'MIN_AFTER_FIRST', 'MAX_AFTER_FIRST', 'IMIN', 'IMAX', 'IAVG', 'TEXT')
LEARN_HEADERS = ('BATCH_SIZE', 'MFG_MONTH', 'EXP_MONTH', 'BMR_NO', 'MATERIAL_CODE', 'PRODUCT_NAME',
                 'MANUFACTURING_ORDER', 'MFG_START')
LEARN_MIN_RATE = 0.8


def _close(a, b):
    return abs(a - b) <= max(0.011, abs(b) * 0.001)


def _theme(wb):
    xml = wb.loaded_theme.decode('utf-8', 'replace') if isinstance(wb.loaded_theme, bytes) else (wb.loaded_theme or '')
    out = []
    for n in ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6']:
        m = re.search(r'<a:%s>.*?(?:srgbClr val="([0-9A-Fa-f]{6})"|lastClr="([0-9A-Fa-f]{6})").*?</a:%s>' % (n, n), xml, re.S)
        out.append((m.group(1) or m.group(2)).upper() if m else 'FFFFFF')
    return out


def _fill(cell, theme):
    import colorsys
    f = cell.fill
    if f is None or f.fill_type != 'solid' or f.fgColor is None:
        return None
    c = f.fgColor
    if c.type == 'rgb' and isinstance(c.rgb, str):
        return c.rgb[-6:].upper()
    if c.type == 'theme' and c.theme is not None and c.theme < len(theme):
        r, g, b = (int(theme[c.theme][i:i + 2], 16) / 255 for i in (0, 2, 4))
        h, l, sat = colorsys.rgb_to_hls(r, g, b)
        t = c.tint or 0
        l = l * (1 + t) if t < 0 else l * (1 - t) + t
        return ''.join('%02X' % round(x * 255) for x in colorsys.hls_to_rgb(h, l, sat))
    return None


def qa_sheet(wb, ws, known, theme):
    """Layout of one sheet of QA's workbook: dict like LAYOUTS + rows {batch: row}; None when it holds no batches."""
    best = None
    for c in range(1, 7):                                    # the batch column: most batches of the export
        hits = {}
        for (r, cc), cell in ws._cells.items():
            if cc == c and type(cell).__name__ != 'MergedCell' and cell.value is not None:
                b = batch_no(resolved(wb, cell.value))
                if b in known and b not in hits:
                    hits[b] = r
        if hits and (best is None or len(hits) > len(best[1])):
            best = (c, hits)
    if not best:
        return None
    bcol, rows = best
    first = None                                             # first batch-like row: the header block ends above
    for r in range(1, min(rows.values()) + 1):
        v = s(resolved(wb, ws.cell(r, bcol).value))
        if re.fullmatch(r'[A-Z0-9/-]*\d{5,}[A-Z0-9/-]*', batch_no(v)):
            first = r
            break
    first = first or min(rows.values())
    cells, maxc = [], bcol
    for (r, c), cell in sorted(ws._cells.items()):
        if r >= first or type(cell).__name__ == 'MergedCell':
            continue
        f = _fill(cell, theme)
        if cell.value is None and f is None:
            continue
        maxc = max(maxc, c) if cell.value is not None else maxc
        cells.append((cell.coordinate, cell.value, f))
    merges = [str(m) for m in ws.merged_cells.ranges if m.max_row < first]
    maxc = max([maxc] + [m.max_col for m in ws.merged_cells.ranges if m.max_row < first])
    from openpyxl.utils import get_column_letter as cl
    widths = {cl(c): round(ws.column_dimensions[cl(c)].width, 2) for c in range(1, maxc + 1)
              if ws.column_dimensions[cl(c)].width}
    heights = {r: round(ws.row_dimensions[r].height, 2) for r in range(1, first) if ws.row_dimensions[r].height}
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
    cells = [x for x in cells if column_index_from_string(coordinate_from_string(x[0])[0]) <= maxc]
    return dict(name=ws.title, first=first, bcol=bcol, ncols=maxc, cells=cells, merges=merges, widths=widths,
                heights=heights), rows


def _qa_value(v):
    """QA cell -> number, or normalised texts it may equal (dates as QA and PAS-X write them)."""
    import datetime as _dt
    if isinstance(v, (_dt.datetime, _dt.date)):
        return None, {x.lower().replace(' ', '') for x in (v.strftime('%b.%Y'), v.strftime('%d-%m-%Y'), v.strftime('%b-%Y'),
                                                            v.strftime('%m/%Y'), v.strftime('%d-%m-%Y %H:%M'))}
    n = num(v)
    t = _norm(v).replace(' ', '')
    return n, {t} if t else set()


ROMAN = {'i': '1', 'ii': '2', 'iii': '3', 'iv': '4', 'v': '5'}


def _tokens(text):
    """Words of a header or a PAS-X name for breaking ties (LOT-II = LOT2 = Lot 2); never used to find a value."""
    out = set()
    for w in re.findall(r'[a-z]+|\d+', s(text).lower().replace('-', ' ')):
        out.add(ROMAN.get(w, w))
    for a, b in re.findall(r'([a-z]+)\s*-?\s*(i{1,3}|iv|v|\d)\b', s(text).lower()):
        out.add(a + ROMAN.get(b, b))
    return out - {'of', 'the', 'and', 'in', 'to', 'for', 'min', 'max', 'ave', 'avg', 'nlt', 'nmt', 'kg', 'no'}


def learn(qa_path, batches, name='LEARNED'):
    """(layout, rules, report) learned from QA's own workbook: each QA column -> the exact PAS-X source that gives
    QA's values in the batches both have. report rows = [sheet, column, header path, result, rule source, agreed, of]."""
    import openpyxl
    from openpyxl.utils import get_column_letter as cl
    wb = openpyxl.load_workbook(qa_path) if isinstance(qa_path, str) else openpyxl.load_workbook(io.BytesIO(qa_path))
    theme = _theme(wb)
    lname = 'LEARNED-' + name
    RULESET_VERSION.setdefault(lname, "learned from QA's workbook (curator %s) - QA to approve" % VERSION)
    ser = series(batches)
    rows_of = {}
    for b, rows in batches.items():                          # rows of each series per batch
        d = rows_of.setdefault(b, {})
        for r in rows:
            par = s(r.get('PARAMETER_CODE')) or '@' + ' '.join(s(r.get('APQR_PARAMETER')).split())
            d.setdefault((code(r.get('STEP_CODE_PATH')), par), []).append(r)
    # candidate values: (kind, source, statistic, scale) -> {batch: value}
    cand = OrderedDict()
    for fld in LEARN_HEADERS:
        cand[('HEADER', fld, '', 1)] = {b: header(rows, fld)[0] for b, rows in batches.items()}
    for key, f in ser.items():
        for st in LEARN_STATS:
            if st == 'TEXT' and not f['text']:
                continue
            if st != 'TEXT' and not f['numeric']:
                continue
            if st in ('IMIN', 'IMAX', 'IAVG') and not f['units']:
                continue
            if st in ('MIN_AFTER_FIRST', 'MAX_AFTER_FIRST') and not f['repeated']:
                continue
            vals = {}
            for b in batches:
                rs = rows_of[b].get(key)
                if rs:
                    v = statistic(rs, st)[0]
                    if v is not None:
                        vals[b] = v
            if vals:
                for fac in ((1, 1 / 60.0, 60) if st != 'TEXT' else (1,)):
                    cand[('PASX', _source(*key), st, fac)] = {b: (v * fac if fac != 1 else v) for b, v in vals.items()}
    keys = list(cand)
    idx = {}

    def index(b):
        """Candidate values of one batch: sorted numbers (bisect) and normalised texts."""
        if b not in idx:
            nl, td = [], {}
            for i, k in enumerate(keys):
                v = cand[k].get(b)
                if v is None:
                    continue
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    nl.append((float(v), i))
                t = _norm(v).replace(' ', '')
                if t:
                    td.setdefault(t, []).append(i)
            nl.sort()
            idx[b] = (([x for x, _ in nl], [i for _, i in nl]), td)
        return idx[b]
    pref = {st: i for i, st in enumerate(LEARN_STATS)}
    lay, rules, report = [], [], []
    for ws in wb.worksheets:
        got = qa_sheet(wb, ws, set(batches), theme)
        if not got:
            report.append([ws.title, '', '', 'NO_COMMON_BATCH', '', 0, 0])
            continue
        sh, qrows = got
        lay.append(sh)
        paths = layout_paths(sh)
        under = {}                                           # merged group headings (LOT-II ...) over every column
        for m in ws.merged_cells.ranges:
            if m.max_row < sh['first']:
                v = ws.cell(m.min_row, m.min_col).value
                for c in range(m.min_col, m.max_col + 1):
                    under.setdefault(cl(c), []).append(s(v))
        for c in range(sh['bcol'] + 1, sh['ncols'] + 1):
            col = cl(c)
            if col not in paths:
                continue
            path = ' > '.join(paths[col])
            qv, formulas = {}, []
            for b, r in qrows.items():
                v = ws.cell(r, c).value
                if isinstance(v, str) and v.startswith('='):
                    formulas.append(re.sub(r'(?<=[A-Z])%d\b' % r, '{r}', v))
                elif v not in (None, '') and s(v) not in ('-', 'NA', 'N/A', 'na'):
                    qv[b] = v
            rid = 'L-%s-%s' % (re.sub(r'\W+', '', ws.title)[:12], col)
            if formulas and len(set(formulas)) == 1 and len(formulas) >= len(qv):
                rules.append(_rule(rid, sh['name'], col, 'FORMULA', formulas[0], layout=lname, status='LEARNED',
                                   evidence="QA's own formula in %d batch row(s)" % len(formulas)))
                report.append([sh['name'], col, path, 'FORMULA', formulas[0], len(formulas), len(formulas)])
                continue
            if not qv:
                rules.append(_rule(rid, sh['name'], col, 'QA_ENTRY', layout=lname,
                                   evidence='QA leaves it empty in the %d common batch(es)' % len(qrows)))
                report.append([sh['name'], col, path, 'QA_EMPTY', '', 0, 0])
                continue
            parsed = {b: _qa_value(v) for b, v in qv.items()}
            n = len(parsed)
            need = n if n < 3 else max(3, int(LEARN_MIN_RATE * n + 0.9999))
            agree = []
            refs = list(parsed)[:5]                      # sources that match QA in one of 5 batches, then all batches
            pool = set()
            for b in refs:
                (vals_sorted, ids), td = index(b)
                qn, qt = parsed[b]
                if qn is not None:
                    tol = max(0.011, abs(qn) * 0.001)
                    pool.update(ids[bisect.bisect_left(vals_sorted, qn - tol):bisect.bisect_right(vals_sorted, qn + tol)])
                for t in qt:
                    pool.update(td.get(t, ()))
            for i in pool:
                k = keys[i]
                vals = cand[k]
                a = 0
                for b, (qn, qt) in parsed.items():
                    v = vals.get(b)
                    if v is None:
                        continue
                    if qn is not None and isinstance(v, (int, float)) and _close(float(v), qn):
                        a += 1
                    elif qt and _norm(v).replace(' ', '') in qt:
                        a += 1
                if a >= need:
                    agree.append((a, k))
            if not agree:
                rules.append(_rule(rid, sh['name'], col, 'NO_PASX_SOURCE', layout=lname, evidence=
                                   'no PAS-X series reproduces QA\'s %d value(s) (e.g. %s)' % (n, s(list(qv.values())[0])[:40])))
                report.append([sh['name'], col, path, 'NOT_LEARNED', '', 0, n])
                continue
            varies = len({round(qn, 3) if qn is not None else tuple(sorted(qt)) for qn, qt in parsed.values()}) > 1
            if varies:                                   # a constant PAS-X series cannot stand in for QA's varying values
                agree = [(a, k) for a, k in agree if len({clean(v) if isinstance(v, float) else v for b, v in
                                                          cand[k].items() if b in parsed}) > 1]
            if not agree:
                rules.append(_rule(rid, sh['name'], col, 'NO_PASX_SOURCE', layout=lname, evidence=
                                   "only constant PAS-X series match QA's varying values"))
                report.append([sh['name'], col, path, 'NOT_LEARNED', '', 0, n])
                continue
            want_words = _tokens(path) | _tokens(sh['name']) | _tokens(' '.join(under.get(col, [])))

            def fit(k):                                  # name affinity: a tie-breaker between equally good sources
                if k[0] == 'HEADER':
                    return 0
                f = ser.get(tuple(k[1].split('|', 1)), {})
                have = _tokens(' '.join([f.get('sheet', ''), f.get('stage', ''), f.get('step', ''), f.get('param', ''),
                                         k[1].replace('>', ' ')]))
                return len(want_words & have)
            top = max(a for a, _ in agree)
            best = [k for a, k in agree if a == top]
            if len(best) > 1:
                hi = max(fit(k) for k in best)
                best = [k for k in best if fit(k) == hi]
            # different sources that would write different values in the other batches: not decidable from QA's data
            vec = lambda k: tuple((b, clean(cand[k][b]) if isinstance(cand[k][b], float) else cand[k][b])
                                  for b in batches if b in cand[k])
            distinct = {vec(k) for k in best}
            best.sort(key=lambda k: (k[0] != 'HEADER', k[3] != 1, -fit(k), pref.get(k[2], 99), k[1]))
            k = best[0]
            if len(distinct) > 1:
                rules.append(_rule(rid, sh['name'], col, 'NO_PASX_SOURCE', layout=lname, evidence=
                                   'AMBIGUOUS: %d sources match QA (%d/%d) but differ elsewhere: %s' % (
                                       len(best), top, n, '; '.join('%s %s' % (x[1], x[2]) for x in best[:4]))))
                report.append([sh['name'], col, path, 'AMBIGUOUS', ' ; '.join('%s %s' % (x[1], x[2]) for x in best[:6]), top, n])
                continue
            ev = 'agrees with QA in %d of %d common batch(es)%s' % (top, n, '' if n >= 3 else ' - few batches, QA to confirm')
            if k[0] == 'HEADER':
                rules.append(_rule(rid, sh['name'], col, 'HEADER', k[1], layout=lname, status='LEARNED', evidence=ev))
            else:
                rules.append(_rule(rid, sh['name'], col, 'PASX', k[1], k[2], scale='' if k[3] == 1 else repr(k[3]),
                                   layout=lname, status='LEARNED', evidence=ev))
            report.append([sh['name'], col, path, 'LEARNED', '%s %s%s' % (k[1], k[2], '' if k[3] == 1 else ' x%g' % k[3]), top, n])
    if not lay:
        raise ValueError("QA's workbook has none of the export's batches: nothing to learn from (give an export that "
                         "includes batches QA already has)")
    return lay, rules, report


MODE_TEXT = {'APPROVED': 'approved rule set (exact codes, verified against QA)',
             'LEARNED': "learned from QA's own workbook: QA to approve the rules (sheet APQR_Learning)",
             'GENERIC': "generic QA format for a product without approved rules: QA to review before use"}


def quality(trace, vc):
    """Acceptance metrics. Traceability = written cells whose trace names their PAS-X source (or QA's formula / the
    machine rule) / written cells. With QA's workbook: accuracy = equal / (equal + different) over the cells written,
    coverage = QA's values reproduced / QA's values in the rule columns (QA_EMPTY = written where QA has nothing)."""
    written = [t for t in trace if t[8] in ('FILLED', 'FILLED_PROVISIONAL', 'FILLED_AUTO', 'FILLED_LEARNED', 'FORMULA',
                                            'NA_OTHER_MACHINE') and t[7] not in (None, '')]
    traced = [t for t in written if t[TRACE_COLS.index('SOURCE_TYPE')]]
    m = OrderedDict([('cells written', len(written)),
                     ('traceability', '%.1f %%' % (100.0 * len(traced) / len(written)) if written else 'n/a')])
    same, diff, only = vc.get('SAME', 0), vc.get('DIFFERENT', 0), vc.get('ONLY_QA', 0)
    if same + diff:
        m['accuracy vs QA'] = '%.1f %% (%d of %d)' % (100.0 * same / (same + diff), same, same + diff)
        m['coverage of QA values'] = '%.1f %% (%d of %d)' % (100.0 * same / (same + diff + only), same, same + diff + only)
    return m


def unmapped(batches, rules):
    """PAS-X series of the export that no rule uses: candidates for QA's next rule (sheet APQR_Unmapped_PASX)."""
    srcs = [(p, q) for r in rules if r['KIND'] == 'PASX' for p, q in parse_sources(r['SOURCES'])]
    used = set()
    for b in batches.values():
        for p, q in set(srcs):
            for r in match(b, p, q):
                used.add((code(r.get('STEP_CODE_PATH')), s(r.get('PARAMETER_CODE')) or
                          '@' + ' '.join(s(r.get('APQR_PARAMETER')).split())))
    out = []
    for (path, par), f in series(batches).items():
        if (path, par) in used:
            continue
        rows = [r for b in batches.values() for r in b if code(r.get('STEP_CODE_PATH')) == path and
                (s(r.get('PARAMETER_CODE')) or '@' + ' '.join(s(r.get('APQR_PARAMETER')).split())) == par]
        ex = next((s(r.get('READING_VALUES')) or s(r.get('READING_MAX')) or s(r.get('ACTUAL_TEXT')) for r in rows), '')
        out.append([f['sheet'], f['stage'], f['step'], f['param'], f['uom'], path, par,
                    len({batch_no(r.get('BATCH_NO')) for r in rows}), ex[:60]])
    return out


# ------------------------------------------------------------------------------------------------ which rule set
def detect_layout(batches):
    """(MBR rule set or None, reason)."""
    mbrs = sorted({first(b, 'MBR_ID') for b in batches.values()} - {''})
    for m in mbrs:
        if m in RULESETS:
            return m, 'MBR_ID %s in the export' % m
    mats = {first(b, 'MATERIAL_CODE').lstrip('0') for b in batches.values()}
    codes = {code(r.get('STEP_CODE_PATH')) for b in batches.values() for r in b}
    for m, sign in LAYOUT_SIGNS.items():
        if mats & {x.lstrip('0') for x in sign['materials']}:
            return m, 'material %s (MBR_ID %s)' % (', '.join(sorted(mats)), ', '.join(mbrs) or 'empty')
        if sum(c in codes for c in sign['codes']) >= 2:
            return m, 'step codes of MBR %s (MBR_ID %s)' % (m, ', '.join(mbrs) or 'empty')
    return None, 'no approved rule set for MBR %s / material %s' % (', '.join(mbrs) or 'empty', ', '.join(sorted(mats)))


# ------------------------------------------------------------------------------------------------ validation
def likely_cause(qv, cv, rule, b, field=''):
    """Why QA's value differs, from the data alone: a digit slip, QA's value is another PAS-X series of the batch, the
    batch ran on the other machine, or QA's value is nowhere in the export (check the batch record)."""
    q, c = num(qv), num(cv)
    if q is None and isinstance(qv, str) and re.fullmatch(r'[\d.,\s]+', qv.strip()):
        return 'QA_TYPING: %r is not a number' % qv
    if q is not None and c is not None and c != 0:
        for f in (10, 100, 1000, 0.1, 0.01, 0.001):
            if abs(q - c * f) <= max(0.011, abs(q) * 0.001):
                return 'QA_TYPING: decimal point (QA %s, PAS-X %s)' % (clean(q), clean(c))
    if s(cv).upper() == 'NA' and rule.get('MACHINE'):
        return 'MACHINE: PAS-X shows the batch on the %s-station machine only; QA has values in the %s-station block' % (
            station(b), rule['MACHINE'])
    if q is None:
        return 'TEXT: QA %s, PAS-X %s - check which the batch record supports' % (s(qv)[:30], s(cv)[:30])
    own = {(code(path), par) for path, par in parse_sources(rule.get('SOURCES', ''))}
    hits = []
    for r in b:
        key = (code(r.get('STEP_CODE_PATH')), s(r.get('PARAMETER_CODE')))
        if key in own:
            continue
        for f in ('READING_MIN', 'READING_MAX', 'READING_AVG', 'READING_MIN_AFTER_FIRST', 'READING_MAX_AFTER_FIRST'):
            v = num(r.get(f))
            if v is not None and abs(v - q) <= max(0.011, abs(q) * 0.001):
                hits.append('%s|%s (%s)' % (key[0], key[1], s(r.get('APQR_PARAMETER'))))
                break
    if hits:
        want = _tokens(field) | _tokens(rule.get('SOURCES', '').replace('>', ' '))
        hits = sorted(dict.fromkeys(hits), key=lambda h: -len(want & _tokens(h.replace('>', ' '))))
        return 'OTHER_SOURCE: QA\'s value = PAS-X %s%s - the rule uses %s; check which one the QA header means' % (
            hits[0], ' (+%d more)' % (len(hits) - 1) if len(hits) > 1 else '', rule.get('SOURCES', '')[:80])
    return 'NOT_IN_PASX: QA\'s value is not in the export - check the batch record page'


def validate(wb, qa_path, layout, rules, order, batches=None):
    """Curated cells against QA's own workbook for the batches QA already has: SAME / DIFFERENT / QA_EMPTY / ONLY_QA;
    each DIFFERENT with its likely cause (likely_cause)."""
    import openpyxl
    qa = openpyxl.load_workbook(qa_path)
    out = []
    for sh in layout:
        if not sh['bcol'] or sh['name'] not in qa.sheetnames or sh['name'] not in wb.sheetnames:
            continue
        from openpyxl.utils import get_column_letter
        bcol = get_column_letter(sh['bcol'])
        qrows, _, _ = batch_rows(qa, qa[sh['name']], bcol, sh['first'])
        crows, _, _ = batch_rows(wb, wb[sh['name']], bcol, sh['first'])
        for b in order:
            if b not in qrows or b not in crows:
                continue
            for rule in rules:
                if rule['SHEET'] != sh['name'] or rule['KIND'] == 'FORMULA':
                    continue
                qv = qa[sh['name']]['%s%d' % (rule['COLUMN'], qrows[b])].value
                cv = wb[sh['name']]['%s%d' % (rule['COLUMN'], crows[b])].value
                if isinstance(qv, str) and qv.startswith('='):
                    continue
                if qv in (None, '') and cv in (None, ''):
                    continue
                a, q = num(cv), num(qv)
                if qv in (None, ''):
                    k = 'QA_EMPTY'
                elif cv in (None, ''):
                    k = 'ONLY_QA'
                elif a is not None and q is not None:
                    k = 'SAME' if abs(a - q) <= max(0.011, abs(q) * 0.001) else 'DIFFERENT'
                else:
                    k = 'SAME' if _norm(cv).replace(' ', '') == _norm(qv).replace(' ', '') else 'DIFFERENT'
                out.append([b, sh['name'], '%s%d' % (rule['COLUMN'], crows[b]), field_name(sh, rule['COLUMN']), qv, cv,
                            k, rule['RULE_ID'], likely_cause(qv, cv, rule, batches[b], field_name(sh, rule['COLUMN'])) if k == 'DIFFERENT' and batches
                            and b in batches else ''])
    return out


# ------------------------------------------------------------------------------------------------ main entry
def curate(rows, batches_wanted=None, template=None, layout=None, rules=None, validate_with=None, extra_info=(),
           learn_from=None, strict=False):
    """rows = export rows (dicts); template = path / bytes of the approved template workbook (None = built-in header
    copy); layout = force a rule set ('F1M00011'); rules = rule list (default: the built-in approved set);
    validate_with = QA's own workbook to compare with; learn_from = QA's own workbook to learn the layout and rules
    from (a product without an approved rule set); strict = refuse a product without an approved rule set instead of
    the generic QA format. Returns (workbook, trace, exceptions, summary)."""
    import openpyxl
    batches = by_batch(rows)
    if batches_wanted:
        want = [batch_no(x) for x in batches_wanted]
        batches = OrderedDict((b, batches[b]) for b in batches if b in want)
    if not batches:
        raise ValueError('No batches in the export' + (' for the requested batch numbers' if batches_wanted else ''))
    mbrs = sorted({first(b, 'MBR_ID') for b in batches.values()} - {''})
    mats = sorted({first(b, 'MATERIAL_CODE') for b in batches.values()} - {''})
    learned, notes = None, []
    if learn_from is not None:
        name = '_'.join(mbrs) or '_'.join(mats) or 'product'
        lay, lrules, learned = learn(learn_from, by_batch(rows), name)     # all export batches teach, not only the asked
        lrules = no_ave(lrules, lay)
        mbr, mode, why = 'LEARNED-' + name, 'LEARNED', "layout and rules learned from QA's own workbook"
        rules = rules if rules is not None else lrules
    else:
        if layout:
            if layout not in RULESETS:
                raise ValueError('No approved rule set for %r (known: %s)' % (layout, ', '.join(RULESETS)))
            mbr, why = layout, 'chosen with --layout'
        else:
            mbr, why = detect_layout(batches)
        if mbr:
            mode, lay = 'APPROVED', LAYOUTS[mbr]
            rules = rules if rules is not None else RULESETS[mbr]()
        elif strict:
            raise ValueError(why + '. A new product needs its own rule set (rules CSV, approved by QA) and template; '
                             'nothing is filled by guessing.')
        else:
            lay, arules = generic_layout(batches, '_'.join(mbrs) or '_'.join(mats) or 'unknown')
            mbr, mode = arules[0]['LAYOUT'], 'GENERIC'
            why += ": generic QA format (QA's layout, filled where PAS-X has the field's exact code pattern)"
            rules = rules if rules is not None else arules
            if template is not None:
                notes.append('template not used: it belongs to an approved rule set, this product has none')
                template = None
    other = {m for m in mbrs if m != mbr} if mode == 'APPROVED' else set()
    use = OrderedDict((b, r) for b, r in batches.items() if first(r, 'MBR_ID') not in other)
    skipped = [b for b in batches if b not in use]
    version_exc = []
    if mode == 'APPROVED' and mbr in APPROVED_MBR_VERSIONS:
        ok_v = APPROVED_MBR_VERSIONS[mbr]
        for b in list(use):
            v = re.sub(r'\.0+$', '', first(use[b], 'MBR_VERSION')).lstrip('0')
            if v and v not in ok_v:
                del use[b]
                skipped.append(b)
                version_exc.append([b, '-', '-', 'whole batch', 'MBR-VERSION', 'UNKNOWN_MBR_VERSION',
                                    'MBR %s version %s: the rule set is approved for version(s) %s only - batch not filled '
                                    'until QA approves the rules for this version' % (mbr, v, ', '.join(ok_v))])
    if not use:
        raise ValueError('No batch can be filled: ' + '; '.join(e[6] for e in version_exc[:3]))
    order = sorted(use, key=lambda b: (start_key(first(use[b], 'MFG_START')), b))
    from_template = template is not None
    wb = openpyxl.load_workbook(template if isinstance(template, str) else io.BytesIO(template)) if from_template \
        else builtin_workbook(lay)
    trace, exc = [], list(version_exc)
    names = {n.lower().strip(): n for n in wb.sheetnames}
    master_rows = {}
    for sh in lay:
        name = names.get(sh['name'].lower().strip())
        if name is None:
            notes.append('sheet %r not in the template: not filled' % sh['name'])
            continue
        ws = wb[name]
        if from_template:
            same, diff = header_check(sh, ws)
            notes.append('%s: %d of %d header paths identical' % (name, same, same + len(diff)))
            moved = [c for c in diff if any(r['COLUMN'] == c and r['SHEET'] == sh['name'] and r['KIND'] in
                                            ('PASX', 'HEADER', 'FORMULA') for r in rules)]
            if moved:
                raise ValueError('Template sheet %r: headers of column(s) %s differ from the approved layout %s - '
                                 'refusing to fill (values would land in the wrong columns).'
                                 % (name, ', '.join(moved), mbr))
        if sh['name'] == 'Product Information':
            b0 = use[order[0]] if order else []
            for row, field in ((1, 'PRODUCT_NAME'), (4, 'MATERIAL_CODE')):
                if ws.cell(row, 2).value in (None, ''):
                    v = first(b0, field)
                    ws.cell(row, 2).value = int(v) if field == 'MATERIAL_CODE' and v.isdigit() else v
            continue
        if not sh['bcol']:
            continue
        srules = [r for r in rules if r['SHEET'] == sh['name']]
        have = fill_sheet(wb, ws, sh, srules, use, order, trace, exc, master_rows, from_template)
        if sh['name'].upper().startswith('BMR'):
            master_rows = dict(have)
    filled = sum(1 for t in trace if t[8] in ('FILLED', 'FILLED_PROVISIONAL', 'FILLED_AUTO', 'FILLED_LEARNED'))
    counts = OrderedDict()
    for t in trace:
        counts[t[8]] = counts.get(t[8], 0) + 1
    validation = validate(wb, validate_with, lay, rules, order, use) if validate_with else None
    vc = OrderedDict()
    for v in validation or []:
        vc[v[6]] = vc.get(v[6], 0) + 1
    metrics = quality(trace, vc)
    if mode != 'APPROVED':                            # amber tabs: layout or rules not approved by QA yet
        for sh in lay:
            if sh['name'] in wb.sheetnames:
                wb[sh['name']].sheet_properties.tabColor = 'FFC000'
    info = list(extra_info) + [('Tool', 'AdigiPQR curator v' + VERSION), ('Mode', MODE_TEXT[mode]),
            ('Rule set', RULESET_VERSION.get(mbr, mbr)),
            ('Rule set chosen because', why), ('Template', 'given workbook' if from_template else
                                               'built-in copy of QA header rows'),
            ('MBR(s) in export', ', '.join(mbrs) or '?'), ('Batches', len(order)),
            ('Batches of another MBR (left out)', ', '.join(skipped) or 'none')] + \
        [('Cells ' + k, v) for k, v in counts.items()] + [('Exceptions', len(exc))] + \
        [('Validation ' + k, v) for k, v in vc.items()] + [('Metric ' + k, v) for k, v in metrics.items()] + \
        [('Template check', n) for n in notes] + \
        [('Principle', 'Every filled cell comes from an exact PAS-X step code path + parameter code named in a rule; '
                       'blank + exception instead of a guess. Prototype for QA review: verify against the batch record.')]
    own_sheets(wb, trace, exc, rules, info, validation, learned, unmapped(use, rules))
    summ = dict(mode=mode, learned=learned, metrics=metrics, differences=[v for v in validation or [] if v[6] == 'DIFFERENT'], preview=lambda: preview(wb, lay, trace), batches=len(order), order=order, filled=filled, exceptions=len(exc), counts=counts, mbr=mbr, why=why,
                skipped=skipped, notes=notes, validation=vc, template=from_template)
    return wb, trace, exc, summ


def preview(wb, lay, trace, max_rows=400, max_cols=700):
    """The curated sheets as the web page shows them: header rows (merged, coloured) + the rows of this run's batches,
    each written cell with its status and PAS-X source (for QA to verify without opening the file)."""
    import datetime as _dt
    from openpyxl.utils import column_index_from_string as ci
    from openpyxl.utils.cell import coordinate_from_string
    info = {}
    for t in trace:
        col, row = coordinate_from_string(t[5])
        src = ' ; '.join(x for x in (t[11] and '%s | %s' % (t[11], t[12]), t[14], t[20], t[21]) if x)
        info[(t[4], row, ci(col))] = (t[8], src[:400])

    def show(v):
        if v is None:
            return ''
        if isinstance(v, float):
            return str(clean(v))
        if isinstance(v, (_dt.datetime, _dt.date)):
            return v.strftime('%d-%m-%Y')
        return str(v)
    out = []
    for sh in lay:
        if sh['name'] not in wb.sheetnames or not sh.get('bcol'):
            continue
        ws = wb[sh['name']]
        rows = sorted({r for (n, r, _c) in info if n == sh['name']})[:max_rows]
        if not rows:
            continue
        ncol = min(max(sh['ncols'], max(c for (n, _r, c) in info if n == sh['name'])), max_cols)
        fills = {coordinate_from_string(c)[1] * 100000 + ci(coordinate_from_string(c)[0]): f for c, _v, f in sh['cells'] if f}
        head = [[show(ws.cell(r, c).value) if type(ws.cell(r, c)).__name__ != 'MergedCell' else None
                 for c in range(1, ncol + 1)] for r in range(1, sh['first'])]
        merges = []
        for m in ws.merged_cells.ranges:
            if m.max_row < sh['first'] and m.min_col <= ncol:
                merges.append([m.min_row, m.min_col, m.max_row, min(m.max_col, ncol)])
        body = []
        for r in rows:
            line = []
            for c in range(1, ncol + 1):
                st, src = info.get((sh['name'], r, c), ('', ''))
                line.append([show(resolved(wb, ws.cell(r, c).value)), st, src])
            body.append(line)
        out.append(dict(name=sh['name'], head=head, merges=merges, body=body,
                        fills={str(k): v for k, v in fills.items() if k // 100000 < sh['first']}))
    return out


def find_template(paths, mbr):
    """An approved template next to the export or this program: APQR_TEMPLATE_<MBR>.xlsx or '*Blank*Template*.xlsx'."""
    dirs = list(dict.fromkeys([os.path.dirname(os.path.abspath(p)) for p in paths] +
                              [os.path.dirname(os.path.abspath(__file__)), os.getcwd()]))
    for d in dirs:
        for f in sorted(os.listdir(d)):
            low = f.lower()
            if low.endswith('.xlsx') and not low.startswith('~$') and (
                    low == ('apqr_template_%s.xlsx' % mbr).lower() or ('blank' in low and 'template' in low)):
                return os.path.join(d, f)
    return None


def main(argv=None):
    ap = argparse.ArgumentParser(description='PAS-X raw export -> QA APQR workbook (approved rules, exact template).')
    ap.add_argument('exports', nargs='*', help='export .xlsx / .csv of final/APQR_FINAL_QUERY.sql')
    ap.add_argument('--template', help='approved APQR template .xlsx (default: APQR_TEMPLATE_<MBR>.xlsx or a '
                                       '"...Blank...Template...xlsx" next to the export / this program)')
    ap.add_argument('-o', '--out', help='output .xlsx')
    ap.add_argument('--batches', help='only these batch numbers, comma separated')
    ap.add_argument('--layout', help='force a rule set, e.g. F1M00011 (normally found from MBR / material / steps)')
    ap.add_argument('--rules', help='rule registry CSV (QA-edited copy of --export-rules)')
    ap.add_argument('--export-rules', metavar='CSV', help='write the built-in rule registry to CSV and stop')
    ap.add_argument('--validate', metavar='QA_WORKBOOK', help="compare the result with QA's own workbook")
    ap.add_argument('--learn', metavar='QA_WORKBOOK', help="learn layout and rules from QA's own workbook (a product "
                                                           "without an approved rule set; the export must share batches)")
    ap.add_argument('--qa', metavar='QA_WORKBOOK', help="QA's own workbook of the product: --validate when the product has "
                                                       'an approved rule set, else --learn and --validate')
    ap.add_argument('--preview-json', metavar='JSON', help='also write the sheets as the web page shows them')
    ap.add_argument('--strict', action='store_true', help='refuse a product without an approved rule set (no generic QA format)')
    a = ap.parse_args(argv)
    if a.export_rules:
        write_rules(a.export_rules, RULESETS[a.layout or 'F1M00011']())
        print('Rule registry written: ' + os.path.abspath(a.export_rules))
        return
    if not a.exports:
        ap.error('give the export file(s)')
    rows = []
    for p in a.exports:
        try:
            rows += read_export(p)
        except ValueError as e:
            sys.exit('STOPPED: %s' % e)
    if not rows:
        sys.exit('No rows with a BATCH_NO column in %s - is this the export of APQR_FINAL_QUERY.sql?' % ', '.join(a.exports))
    wanted = [x for x in re.split(r'[,\s]+', a.batches or '') if x]
    mbr, _ = (a.layout, '') if a.layout else detect_layout(by_batch(rows))
    if a.qa:
        a.validate = a.validate or a.qa
        if not mbr:
            a.learn = a.learn or a.qa
    template = a.template or (find_template(a.exports, mbr) if mbr and not a.learn else None)
    try:
        wb, trace, exc, summ = curate(rows, wanted or None, template=template, layout=a.layout,
                                      rules=read_rules(a.rules) if a.rules else None, validate_with=a.validate,
                                      learn_from=a.learn, strict=a.strict)
    except ValueError as e:
        sys.exit('STOPPED: %s' % e)
    mat = '_'.join(sorted({s(r.get('MATERIAL_CODE')) for r in rows} - {''})) or 'export'
    out = a.out or 'APQR_%s.xlsx' % mat
    wb.save(out)
    print('AdigiPQR curator %s - %s' % (VERSION, RULESET_VERSION.get(summ['mbr'], summ['mbr'])))
    print('Mode: %s' % MODE_TEXT[summ['mode']])
    print('Rule set: %s  (%s)' % (summ['mbr'], summ['why']))
    print('Template: %s' % (template if summ['template'] else 'built-in copy of QA header rows' if summ['mode'] ==
                            'APPROVED' else "QA's workbook header rows" if summ['mode'] == 'LEARNED' else
                            'generic QA format (built in)'))
    if summ['learned'] is not None:
        lc = OrderedDict()
        for x in summ['learned']:
            lc[x[3]] = lc.get(x[3], 0) + 1
        print("Learned from QA's workbook: " + ', '.join('%s %d' % kv for kv in lc.items()) + ' (sheet %s)' % LEARN_SHEET)
    for n in summ['notes']:
        print('  ' + n)
    print('%d export rows, %d batches: %s' % (len(rows), summ['batches'],
                                               ', '.join('%s %d' % (k, v) for k, v in summ['counts'].items())))
    if 'READING_VALUES' not in rows[0]:
        print('WARNING: this export was made before SQL v6 (no repetitions): Min / Max / Avg cells are left blank.')
    if summ['exceptions']:
        print('%d exception(s) - see sheet %s' % (summ['exceptions'], EXC_SHEET))
    if summ['validation']:
        print("Against QA's workbook: " + ', '.join('%s %d' % (k, v) for k, v in summ['validation'].items()))
        causes = OrderedDict()
        for d in summ['differences']:
            k = d[8].split(':')[0]
            causes[k] = causes.get(k, 0) + 1
        if causes:
            print('  differences by likely cause: ' + ', '.join('%s %d' % kv for kv in causes.items()) +
                  ' (sheet %s, column LIKELY_CAUSE)' % VALID_SHEET)
    if summ['skipped']:
        print('Batches left out (another MBR or an MBR version without approved rules): ' + ', '.join(summ['skipped']))
    print('Metrics: ' + '; '.join('%s %s' % kv for kv in summ['metrics'].items()))
    print('Saved: ' + os.path.abspath(out))
    if a.preview_json:
        import json
        with open(a.preview_json, 'w', encoding='utf-8') as fh:
            json.dump(dict(sheets=summ['preview'](), mode=summ['mode'], counts=summ['counts'], batches=summ['order'],
                           metrics=summ['metrics'], skipped=summ['skipped'], version=VERSION,
                           differences=[[s(x) for x in d] for d in summ['differences'][:500]],
                           exceptions=summ['exceptions'], validation=summ['validation']), fh)


# ------------------------------------------------------------------------------------------------ QA layouts
# Header rows of QA's own workbook (tools/make_layout.py; headers only, never a batch row): sheet name, first
# batch row, batch column, number of columns, header cells (cell, text, fill RGB), merged ranges, widths, heights.
LAYOUTS = {'F1M00011': [{'name': 'Product Information',
               'first': 6,
               'bcol': None,
               'ncols': 2,
               'cells': [('A1', 'Name', None), ('A2', 'Label Claim', None), ('A3', 'Storage condition', None),
                         ('A4', 'Product code No. ', None), ('A5', 'Description', None)],
               'merges': [],
               'widths': {'A': 19.86, 'B': 73.71},
               'heights': {2: 30.0, 3: 31.5, 5: 15.75}},
              {'name': 'BMR DETAILS',
               'first': 3,
               'bcol': 3,
               'ncols': 15,
               'cells': [('A1', None, '8EB4E3'), ('B1', None, '8EB4E3'), ('C1', None, '8EB4E3'),
                         ('D1', None, '8EB4E3'), ('E1', None, '8EB4E3'), ('A2', 'Sr. No.', 'DDD9C3'),
                         ('B2', 'Sr. No.', 'DDD9C3'), ('C2', 'BMR No.', 'DDD9C3'),
                         ('D2', 'Batch Size (In Kg)', '00B050'), ('E2', 'Mfg. Date', 'DBEEF4'),
                         ('F2', 'Exp. Date', 'DBEEF4'), ('G2', 'BMR No.', 'DBEEF4'), ('H2', 'Ref. Doc No.', 'DBEEF4'),
                         ('I2', 'SFG Product Code', 'DBEEF4'), ('J2', 'Market', 'DBEEF4'),
                         ('K2', 'OOS/OOT/ Deviation/ Change control', 'DBEEF4'), ('L2', 'Remarks', 'DDD9C3'),
                         ('M2', 'Trending By', 'DDD9C3'), ('N2', 'Estabilished control limit', 'DDD9C3'),
                         ('O2', 'Verified By', 'DDD9C3')],
               'merges': ['F1:L1'],
               'widths': {'A': 4.86,
                          'B': 13.0,
                          'C': 14.14,
                          'D': 9.86,
                          'E': 10.57,
                          'F': 9.14,
                          'G': 20.43,
                          'H': 15.14,
                          'I': 10.29,
                          'J': 9.71,
                          'K': 23.57,
                          'L': 17.43,
                          'M': 18.43,
                          'N': 11.57,
                          'O': 15.29},
               'heights': {1: 24.75, 2: 51.75}},
              {'name': 'Granulation batch size 297.52',
               'first': 7,
               'bcol': 3,
               'ncols': 55,
               'cells': [('A1', 'Sr. No.', 'DDD9C3'), ('B1', 'Sr. No.', 'DDD9C3'),
                         ('C1', 'Manufacturing Batch No. ↓', 'DDD9C3'), ('D1', 'Equipment ID', 'DDD9C3'),
                         ('E1', 'Equipment ID', 'DDD9C3'), ('F1', None, 'DDD9C3'),
                         ('G1', 'Granulation LOT-I', 'DBEEF4'), ('AD1', 'Granulation LOT-II', 'DBEEF4'),
                         ('BA1', 'Remark', 'DDD9C3'), ('BB1', None, 'DBEEF4'), ('BC1', 'Remark', 'DDD9C3'),
                         ('C2', 'Stage', 'DDD9C3'), ('F2', None, 'DDD9C3'), ('G2', 'Wet granulation', 'DBEEF4'),
                         ('R2', 'Wet granulation', 'DBEEF4'), ('S2', 'Drying', 'DBEEF4'),
                         ('AD2', 'Wet granulation', 'DBEEF4'), ('AP2', 'Drying', 'DBEEF4'), ('AW2', None, 'DBEEF4'),
                         ('BB2', 'Blending\n% LOD', '92D050'), ('C3', 'Parameters', 'DDD9C3'), ('F3', None, 'DDD9C3'),
                         ('G3', 'Wet granulation (Binder addition)', 'DBEEF4'),
                         ('J3', 'Wet granulation (Kneading Time/End Point detection)', 'DBEEF4'),
                         ('M3', None, 'DBEEF4'), ('N3', None, 'DBEEF4'), ('O3', 'Ethanol anhydrous ', 'DBEEF4'),
                         ('S3', 'Inlet Temperature ', 'DBEEF4'), ('U3', 'Product Temp.', 'DBEEF4'),
                         ('W3', 'Exhaust Temp', 'FF0000'), ('Y3', '% Flap', 'DBEEF4'),
                         ('AA3', 'Drying Time', 'DBEEF4'), ('AB3', '% LOD', '92D050'),
                         ('AC3', 'Milled Granules % LOD', 'DBEEF4'),
                         ('AD3', 'Wet granulation (Binder addition)', 'DBEEF4'),
                         ('AG3', 'Wet granulation (Kneading Time/End Point detection)', 'DBEEF4'),
                         ('AK3', None, 'DBEEF4'), ('AL3', 'Ethanol anhydrous ', 'DBEEF4'),
                         ('AP3', 'Inlet Temperature ', 'DBEEF4'), ('AR3', 'Product Temp.', 'DBEEF4'),
                         ('AT3', 'Exhaust Temp', 'FF0000'), ('AV3', '% Flap', 'DBEEF4'),
                         ('AX3', 'Drying Time', 'DBEEF4'), ('AY3', '% LOD', '92D050'),
                         ('AZ3', 'Milled Granules % LOD', 'DBEEF4'), ('F4', None, 'DDD9C3'),
                         ('G4', 'Binder addition Time', 'DBEEF4'), ('H4', 'Impeller Speed', 'DBEEF4'),
                         ('I4', 'Chopper Speed', 'DBEEF4'), ('J4', 'Mixing Time (After complete addition)', 'DBEEF4'),
                         ('K4', 'Impeller Speed', 'DBEEF4'), ('L4', 'Chopper Speed', 'DBEEF4'),
                         ('M4', 'Impeller Load at the end of granulation ', 'DBEEF4'),
                         ('N4', 'Quantity of Ethanol anhydrous ', 'DBEEF4'),
                         ('O4', 'Mixing Time (After complete addition)', 'DBEEF4'),
                         ('P4', 'Impeller Speed', 'DBEEF4'), ('Q4', 'Chopper Speed', 'DBEEF4'),
                         ('R4', 'Impeller Load at the end of granulation ', 'DBEEF4'),
                         ('AD4', 'Binder addition Time', 'DBEEF4'), ('AE4', 'Impeller Speed', 'DBEEF4'),
                         ('AF4', 'Chopper Speed', 'DBEEF4'),
                         ('AG4', 'Mixing Time (After complete addition)', 'DBEEF4'), ('AH4', None, 'DBEEF4'),
                         ('AI4', 'Chopper Speed', 'DBEEF4'),
                         ('AJ4', 'Impeller Load at the end of granulation ', 'DBEEF4'),
                         ('AK4', 'Quantity of Ethanol anhydrous ', 'DBEEF4'),
                         ('AL4', 'Mixing Time (After complete addition)', 'DBEEF4'),
                         ('AM4', 'Impeller Speed', 'DBEEF4'), ('AN4', 'Chopper Speed', 'DBEEF4'),
                         ('AO4', 'Impeller Load at the end of granulation ', 'DBEEF4'), ('C5', 'Limits', 'DDD9C3'),
                         ('F5', None, 'DDD9C3'), ('G5', '180 sec', 'F2DCDB'), ('H5', 'Slow', 'F2DCDB'),
                         ('I5', 'off', 'F2DCDB'), ('J5', '120 to 180 Sec', 'F2DCDB'), ('K5', 'Fast', 'F2DCDB'),
                         ('L5', 'Fast', 'F2DCDB'), ('M5', '(23 to 27 AMP)', 'F2DCDB'), ('N5', None, 'F2DCDB'),
                         ('O5', '120 to 180 Sec', 'F2DCDB'), ('P5', 'Fast', 'F2DCDB'), ('Q5', 'Fast', 'F2DCDB'),
                         ('R5', '(23 to 27 AMP)', 'F2DCDB'), ('S5', 'NMT 45°C\n (approx. 35 - 45°C)', 'F2DCDB'),
                         ('U5', 'To be recorded\n(°C)', 'F2DCDB'), ('W5', 'NMT 45°C\n(approx. 30 - 45°C)', 'FF0000'),
                         ('Y5', 'To be recorded\n(%)', 'F2DCDB'),
                         ('AA5', 'Till LOD is NMT 3.00%w/w\n(Min)', 'F2DCDB'),
                         ('AB5', 'Drying NMT 3.00%w/w', 'F2DCDB'), ('AC5', 'NMT 3.00%w/w', 'F2DCDB'),
                         ('AD5', '180 sec', 'F2DCDB'), ('AE5', 'Slow', 'F2DCDB'), ('AF5', 'off', 'F2DCDB'),
                         ('AG5', '120 to 180 sec ', 'F2DCDB'), ('AH5', 'Fast', 'F2DCDB'), ('AI5', 'Fast', 'F2DCDB'),
                         ('AJ5', '\n(23to 27AMP)', 'F2DCDB'), ('AK5', None, 'F2DCDB'),
                         ('AL5', '120 to 180 Sec', 'F2DCDB'), ('AM5', 'Fast', 'F2DCDB'), ('AN5', 'Fast', 'F2DCDB'),
                         ('AO5', '(23 to 27 AMP)', 'F2DCDB'), ('AP5', 'NMT 45°C \n(approx. 35 - 45°C)', 'F2DCDB'),
                         ('AR5', 'To be recorded\n(°C)', 'F2DCDB'),
                         ('AT5', 'NMT 45°C\n(approx. 30 - 45°C)', 'FF0000'), ('AV5', 'To be recorded\n(%)', 'F2DCDB'),
                         ('AX5', 'Drying Till LOD is NMT 3.00%w/w\n(Min)', 'F2DCDB'),
                         ('AY5', 'NMT 3.00%w/w', 'F2DCDB'), ('AZ5', 'NMT 3.00%w/w', 'F2DCDB'),
                         ('BB5', 'NMT 3.00%w/w', 'F2DCDB'), ('D6', 'Dispensing', 'DDD9C3'),
                         ('E6', 'Granulation', 'DDD9C3'), ('F6', 'Blending', 'DDD9C3'), ('N6', None, 'F2DCDB'),
                         ('S6', 'Min', 'F2DCDB'), ('T6', 'Max', 'F2DCDB'), ('U6', 'Min', 'F2DCDB'),
                         ('V6', 'Max', 'F2DCDB'), ('W6', 'Min', 'FF0000'), ('X6', 'Max', 'FF0000'),
                         ('Y6', 'Min', 'F2DCDB'), ('Z6', 'Max', 'F2DCDB'), ('AK6', None, 'F2DCDB'),
                         ('AP6', 'Min', 'F2DCDB'), ('AQ6', 'Max', 'F2DCDB'), ('AR6', 'Min', 'F2DCDB'),
                         ('AS6', 'Max', 'F2DCDB'), ('AT6', 'Min', 'FF0000'), ('AU6', 'Max', 'FF0000'),
                         ('AV6', 'Min', 'F2DCDB'), ('AW6', 'Max', 'F2DCDB')],
               'merges': ['A1:A6', 'AA3:AA4', 'AA5:AA6', 'AB3:AB4', 'AB5:AB6', 'AC3:AC4', 'AC5:AC6', 'AD1:AZ1',
                          'AD2:AO2', 'AD3:AF3', 'AD5:AD6', 'AE5:AE6', 'AF5:AF6', 'AG3:AJ3', 'AG5:AG6', 'AH5:AH6',
                          'AI5:AI6', 'AJ5:AJ6', 'AL3:AO3', 'AL5:AL6', 'AM5:AM6', 'AN5:AN6', 'AO5:AO6', 'AP2:AV2',
                          'AP3:AQ4', 'AP5:AQ5', 'AR3:AS4', 'AR5:AS5', 'AT3:AU4', 'AT5:AU5', 'AV3:AW4', 'AV5:AW5',
                          'AW2:AZ2', 'AX3:AX4', 'AX5:AX6', 'AY3:AY4', 'AY5:AY6', 'AZ3:AZ4', 'AZ5:AZ6', 'B1:B6',
                          'BA1:BA6', 'BB2:BB4', 'BB5:BB6', 'BC1:BC6', 'C3:C4', 'C5:C6', 'D1:D5', 'E1:E5', 'G1:AC1',
                          'G2:Q2', 'G3:I3', 'G5:G6', 'H5:H6', 'I5:I6', 'J3:L3', 'J5:J6', 'K5:K6', 'L5:L6', 'M5:M6',
                          'O3:Q3', 'O5:O6', 'P5:P6', 'Q5:Q6', 'R2:R3', 'R5:R6', 'S2:AC2', 'S3:T4', 'S5:T5', 'U3:V4',
                          'U5:V5', 'W3:X4', 'W5:X5', 'Y3:Z4', 'Y5:Z5'],
               'widths': {'A': 7.29,
                          'B': 13.0,
                          'C': 11.86,
                          'D': 0.14,
                          'E': 13.86,
                          'F': 13.29,
                          'G': 8.43,
                          'H': 8.57,
                          'I': 9.14,
                          'J': 13.14,
                          'K': 9.71,
                          'L': 10.86,
                          'M': 17.43,
                          'N': 12.86,
                          'O': 13.14,
                          'P': 10.57,
                          'Q': 18.71,
                          'R': 18.14,
                          'S': 9.14,
                          'T': 13.0,
                          'U': 13.0,
                          'V': 13.0,
                          'W': 9.14,
                          'X': 13.0,
                          'Y': 9.14,
                          'Z': 6.43,
                          'AA': 14.86,
                          'AB': 8.29,
                          'AC': 14.29,
                          'AD': 11.0,
                          'AE': 7.71,
                          'AF': 8.43,
                          'AG': 8.57,
                          'AH': 9.14,
                          'AI': 10.86,
                          'AJ': 11.57,
                          'AK': 7.86,
                          'AL': 10.0,
                          'AM': 12.71,
                          'AN': 7.57,
                          'AO': 11.14,
                          'AP': 8.43,
                          'AQ': 13.0,
                          'AR': 8.71,
                          'AS': 5.71,
                          'AT': 7.86,
                          'AU': 8.86,
                          'AV': 9.86,
                          'AW': 9.71,
                          'AX': 12.14,
                          'AY': 10.57,
                          'AZ': 20.86,
                          'BA': 9.29,
                          'BB': 12.29,
                          'BC': 14.29},
               'heights': {1: 25.5, 2: 23.25, 3: 27.0, 4: 39.0, 5: 36.0, 6: 18.75}},
              {'name': 'Compression',
               'first': 5,
               'bcol': 3,
               'ncols': 42,
               'cells': [('A1', 'Sr. No.', 'DDD9C3'), ('B1', 'Sr. No.', 'DDD9C3'),
                         ('C1', 'Mini Tablet Batch No. ↓', 'DDD9C3'), ('D1', 'Equipment ID', 'DDD9C3'),
                         ('E1', 'Compression', 'DDD9C3'), ('X1', 'Remark', 'DDD9C3'),
                         ('Y1', 'Compression Parameter(75 station)PD/08/COM05', 'DBEEF4'),
                         ('AG1', 'Remark', 'DDD9C3'),
                         ('AH1', 'Compression Parameter( 73 station)PD/08/COM08, PD/08/COM15', 'DBEEF4'),
                         ('AL1', None, 'DBEEF4'), ('AM1', None, 'DBEEF4'), ('AN1', None, 'DBEEF4'),
                         ('AO1', None, 'DBEEF4'), ('AP1', 'Remark', 'DDD9C3'), ('C2', 'Test→', 'DDD9C3'),
                         ('E2', 'Uniformity of weight', 'DBEEF4'), ('H2', 'Thickness', 'DBEEF4'),
                         ('K2', 'Diameter', 'DBEEF4'), ('N2', 'Tablet Breaking Force (Hardness)', 'DBEEF4'),
                         ('Q2', 'Friability', 'DBEEF4'), ('S2', 'Mass of 30 tablets (M)', 'DBEEF4'),
                         ('V2', 'Average mass of 30 tablets (M/30)', 'DBEEF4'), ('Y2', 'Turret speed', 'DBEEF4'),
                         ('AA2', 'Feeder speed', 'DBEEF4'), ('AC2', 'Pre compression force', 'DBEEF4'),
                         ('AE2', 'Main compression force', 'DBEEF4'), ('AH2', 'Turret speed', 'DBEEF4'),
                         ('AJ2', 'Feeder speed', 'DBEEF4'), ('AL2', 'Pre compression force', 'DBEEF4'),
                         ('AN2', 'Main compression force', 'DBEEF4'), ('C3', 'Limits→', 'DDD9C3'),
                         ('E3', 'Average Mass(mg) ± 10%', 'F2DCDB'),
                         ('H3', '1.90 ± 0.40 mm\n1.50 to 2.30 mm', 'F2DCDB'),
                         ('K3', '4.00 ± 0.20 mm\n3.80 to 4.20 mm', 'F2DCDB'), ('N3', '10 N to 60 N\n', 'F2DCDB'),
                         ('Q3', 'NMT 0.80 %w/w', 'F2DCDB'), ('S3', '669.75 to 740.25 mg', 'F2DCDB'),
                         ('V3', '23.5 mg ± 5 %\n(22.32 to 24.67)', 'F2DCDB'), ('Y3', '\n12 to 30 RPM', 'F2DCDB'),
                         ('AA3', '\n\n20 to 35 RPM\n', 'F2DCDB'), ('AC3', 'NA', 'F2DCDB'), ('AE3', 'KN', 'F2DCDB'),
                         ('AH3', '\n10 to 25 RPM', 'F2DCDB'), ('AJ3', '\n\n08 to 18 RPM\n', 'F2DCDB'),
                         ('AL3', 'NA', 'F2DCDB'), ('AN3', 'KN(NA)', 'F2DCDB'), ('E4', 'min', 'F2DCDB'),
                         ('F4', 'max', 'F2DCDB'), ('G4', 'Ave', 'F2DCDB'), ('H4', 'min', 'F2DCDB'),
                         ('I4', 'max', 'F2DCDB'), ('J4', 'Ave', 'F2DCDB'), ('K4', 'min', 'F2DCDB'),
                         ('L4', 'max', 'F2DCDB'), ('M4', 'Ave', '86E341'), ('N4', 'min', 'F2DCDB'),
                         ('O4', 'max', 'F2DCDB'), ('P4', 'Ave', '86E341'), ('Q4', 'min', 'F2DCDB'),
                         ('R4', 'max', 'F2DCDB'), ('S4', 'min', 'F2DCDB'), ('T4', 'max', 'F2DCDB'),
                         ('U4', 'Ave', '86E341'), ('V4', 'min', 'F2DCDB'), ('W4', 'max', 'F2DCDB'),
                         ('Y4', 'min', 'F2DCDB'), ('Z4', 'max', 'F2DCDB'), ('AA4', 'min', 'F2DCDB'),
                         ('AB4', 'max', 'F2DCDB'), ('AC4', 'min', 'F2DCDB'), ('AD4', 'max', 'F2DCDB'),
                         ('AE4', 'min', 'F2DCDB'), ('AF4', 'max', 'F2DCDB'), ('AH4', 'min', 'F2DCDB'),
                         ('AI4', 'max', 'F2DCDB'), ('AJ4', 'min', 'F2DCDB'), ('AK4', 'max', 'F2DCDB'),
                         ('AL4', 'min', 'F2DCDB'), ('AM4', 'max', 'F2DCDB'), ('AN4', 'min', 'F2DCDB'),
                         ('AO4', 'max', 'F2DCDB')],
               'merges': ['A1:A4', 'AA2:AB2', 'AA3:AB3', 'AC2:AD2', 'AC3:AD3', 'AE2:AF2', 'AE3:AF3', 'AG1:AG4',
                          'AH1:AK1', 'AH2:AI2', 'AH3:AI3', 'AJ2:AK2', 'AJ3:AK3', 'AL2:AM2', 'AL3:AM3', 'AN2:AO2',
                          'AN3:AO3', 'AP1:AP4', 'B1:B4', 'C3:C4', 'D1:D4', 'E1:W1', 'E2:G2', 'E3:G3', 'H2:J2',
                          'H3:J3', 'K2:M2', 'K3:M3', 'N2:P2', 'N3:P3', 'Q2:R2', 'Q3:R3', 'S2:U2', 'S3:U3', 'V2:W2',
                          'V3:W3', 'X1:X4', 'Y1:AF1', 'Y2:Z2', 'Y3:Z3'],
               'widths': {'A': 6.14,
                          'B': 13.0,
                          'C': 13.29,
                          'D': 20.57,
                          'E': 7.71,
                          'F': 9.14,
                          'G': 8.29,
                          'H': 9.14,
                          'I': 13.0,
                          'J': 13.0,
                          'K': 13.0,
                          'L': 13.0,
                          'M': 13.0,
                          'N': 13.0,
                          'O': 9.0,
                          'P': 9.14,
                          'Q': 8.71,
                          'R': 10.57,
                          'S': 6.14,
                          'T': 13.0,
                          'U': 13.0,
                          'V': 10.57,
                          'W': 13.0,
                          'X': 13.0,
                          'Y': 9.14,
                          'Z': 8.0,
                          'AA': 9.14,
                          'AB': 7.43,
                          'AC': 9.14,
                          'AD': 13.0,
                          'AE': 13.0,
                          'AF': 13.0,
                          'AG': 14.43,
                          'AH': 9.14,
                          'AI': 13.0,
                          'AJ': 16.57,
                          'AK': 29.14,
                          'AL': 9.14,
                          'AM': 13.0,
                          'AN': 13.0,
                          'AO': 13.0,
                          'AP': 13.29},
               'heights': {1: 22.5, 2: 46.5, 3: 42.0, 4: 42.0}},
              {'name': 'Coating',
               'first': 6,
               'bcol': 3,
               'ncols': 34,
               'cells': [('A1', 'Sr. No.', 'DDD9C3'), ('B1', 'Sr. No.', 'DDD9C3'),
                         ('C1', 'Mini Tablet Batch No. ↓', 'DDD9C3'), ('D1', 'Equipment ID', 'DDD9C3'),
                         ('E1', 'Equipment ID', 'DDD9C3'), ('F1', 'Coating', 'DBEEF4'), ('K1', 'Remarks', 'DDD9C3'),
                         ('L1', None, 'DDD9C3'), ('M1', 'Spraying', 'B7DEE8'), ('AD1', 'Drying', 'DBEEF4'),
                         ('AH1', 'Remarks', 'DDD9C3'), ('C2', 'Test→', 'DDD9C3'),
                         ('F2', 'Uniformity of Mass', 'DBEEF4'), ('I2', 'Average Weight ', 'DBEEF4'),
                         ('L2', '(Coating solution preparation)Mixing Time', 'DDD9C3'),
                         ('M2', 'Inlet air temperature', 'DBEEF4'), ('O2', 'Exhaust air temperature ', 'FF0000'),
                         ('Q2', 'Bed Temp. ', 'DBEEF4'), ('S2', 'Pan RPM', 'DBEEF4'), ('U2', 'PP RPM', 'FF0000'),
                         ('W2', 'Inlet CFM', 'DBEEF4'), ('Y2', 'Atomizing pressure', 'DBEEF4'),
                         ('AA2', 'Spray rate (g/min)', 'DBEEF4'), ('AC2', '%weight gain', 'DBEEF4'),
                         ('AD2', 'Inlet air temperature', 'DBEEF4'), ('AF2', 'Bed Temp. ', 'DBEEF4'),
                         ('C4', 'Limits→', 'DDD9C3'), ('D4', 'Dispensing', 'DDD9C3'), ('E4', 'Coating', 'DDD9C3'),
                         ('F4', 'Average Mass(mg) ± 10%', 'F2DCDB'),
                         ('I4', '743.7 mg ± 5%\n(706.5 to 780.9)', 'F2DCDB'), ('M4', 'NMT 60°C', 'F2DCDB'),
                         ('O4', 'NMT 50°C', 'FF0000'), ('Q4', '35°C to 45°C', 'F2DCDB'),
                         ('S4', '1 to 5 rpm', 'F2DCDB'),
                         ('U4', 'As required to achieve desired spray rate', 'FF0000'),
                         ('W4', '1400 – 1900', 'F2DCDB'), ('Y4', '\n0.7 - 2.0 bar\n', 'F2DCDB'),
                         ('AA4', '\n400 ±200  g/min', 'F2DCDB'), ('AC4', 'about 5.5 %w/w', 'F2DCDB'),
                         ('AD4', 'NMT 60°C', 'F2DCDB'), ('AF4', '35°C to 45°C', 'F2DCDB'), ('F5', 'Min', 'F2DCDB'),
                         ('G5', 'Max', 'F2DCDB'), ('H5', 'Ave', 'F2DCDB'), ('I5', 'Min', 'F2DCDB'),
                         ('J5', 'Max', 'F2DCDB'), ('L5', '60 min', 'DDD9C3'), ('M5', 'Min', 'F2DCDB'),
                         ('N5', 'Max', 'F2DCDB'), ('O5', 'Min', 'FF0000'), ('P5', 'Max', 'FF0000'),
                         ('Q5', 'Min', 'F2DCDB'), ('R5', 'Max', 'F2DCDB'), ('S5', 'Min', 'F2DCDB'),
                         ('T5', 'Max', 'F2DCDB'), ('U5', 'Min', 'FF0000'), ('V5', 'Max', 'FF0000'),
                         ('W5', 'Min', 'F2DCDB'), ('X5', 'Max', 'F2DCDB'), ('Y5', 'Min', 'F2DCDB'),
                         ('Z5', 'Max', 'F2DCDB'), ('AA5', 'Min', 'F2DCDB'), ('AB5', 'Max', 'F2DCDB'),
                         ('AD5', 'Min', 'F2DCDB'), ('AE5', 'Max', 'F2DCDB'), ('AF5', 'Min', 'F2DCDB'),
                         ('AG5', 'Max', 'F2DCDB')],
               'merges': ['A1:A5', 'AA2:AB3', 'AA4:AB4', 'AC2:AC3', 'AC4:AC5', 'AD1:AG1', 'AD2:AE3', 'AD4:AE4',
                          'AF2:AG3', 'AF4:AG4', 'AH1:AH5', 'B1:B5', 'C2:C3', 'C4:C5', 'D1:D3', 'D4:D5', 'E1:E3',
                          'E4:E5', 'F1:J1', 'F2:H3', 'F4:H4', 'I2:J3', 'I4:J4', 'K1:K5', 'L2:L4', 'M1:AC1', 'M2:N3',
                          'M4:N4', 'O2:P3', 'O4:P4', 'Q2:R3', 'Q4:R4', 'S2:T3', 'S4:T4', 'U2:V3', 'U4:V4', 'W2:X3',
                          'W4:X4', 'Y2:Z3', 'Y4:Z4'],
               'widths': {'A': 4.29,
                          'B': 13.0,
                          'C': 12.86,
                          'D': 12.86,
                          'E': 11.43,
                          'F': 6.14,
                          'G': 5.29,
                          'H': 6.14,
                          'I': 13.0,
                          'J': 13.0,
                          'K': 7.43,
                          'L': 34.29,
                          'M': 8.0,
                          'N': 7.0,
                          'O': 7.43,
                          'P': 7.71,
                          'Q': 6.86,
                          'R': 7.29,
                          'S': 13.0,
                          'T': 7.0,
                          'U': 7.14,
                          'V': 6.57,
                          'W': 8.0,
                          'X': 8.14,
                          'Y': 7.57,
                          'Z': 7.86,
                          'AA': 7.71,
                          'AB': 6.71,
                          'AC': 9.43,
                          'AD': 8.0,
                          'AE': 7.0,
                          'AF': 6.86,
                          'AG': 7.29,
                          'AH': 17.29},
               'heights': {1: 25.5, 2: 15.0, 3: 22.5, 4: 24.75}},
              {'name': 'Yield',
               'first': 4,
               'bcol': 2,
               'ncols': 9,
               'cells': [('A1', 'Sr. No.', 'DDD9C3'), ('B1', 'Mini Tablet Batch No. ↓', 'DDD9C3'),
                         ('C1', 'Yield Parameter ', 'DBEEF4'), ('I1', 'Remarks', 'CCC1DA'),
                         ('B2', 'Stage →', 'DDD9C3'), ('C2', 'Lubrication', 'DBEEF4'),
                         ('D2', 'LB Rejection', 'DBEEF4'), ('E2', 'Compression', 'DBEEF4'),
                         ('F2', 'Compression Rejection', 'DBEEF4'), ('G2', 'Coating', 'DBEEF4'),
                         ('H2', 'Coating Rejection', 'DBEEF4'), ('B3', 'Limits →', 'DDD9C3'),
                         ('C3', '97.50 to 100', 'F2DCDB'), ('D3', 'NA', 'F2DCDB'),
                         ('E3', '96.00 to 100.00', 'F2DCDB'), ('F3', 'NA', 'F2DCDB'), ('G3', '94.00 to100', 'F2DCDB'),
                         ('H3', 'NA', 'F2DCDB')],
               'merges': ['A1:A3', 'C1:H1', 'I1:I3'],
               'widths': {'A': 4.14,
                          'B': 12.43,
                          'C': 12.29,
                          'D': 13.0,
                          'E': 12.14,
                          'F': 13.0,
                          'G': 11.86,
                          'H': 13.0,
                          'I': 10.43},
               'heights': {1: 25.5, 2: 25.5}},
              {'name': 'Temp & RH',
               'first': 4,
               'bcol': 2,
               'ncols': 13,
               'cells': [('A1', 'Sr. No.', 'DDD9C3'), ('B1', 'Manufacturing Batch No. ↓', 'DDD9C3'),
                         ('C1', 'GRAN', 'DBEEF4'), ('E1', 'BLEND', 'DBEEF4'), ('G1', 'COMP', 'DBEEF4'),
                         ('I1', 'Coating', 'DBEEF4'), ('K1', 'Dispensing', 'DBEEF4'), ('M1', 'Remark', 'CCC1DA'),
                         ('B2', 'Stage →', 'DDD9C3'), ('C2', 'TEMP', 'DBEEF4'), ('D2', 'RH', 'DBEEF4'),
                         ('E2', 'TEMP', 'DBEEF4'), ('F2', 'RH', 'DBEEF4'), ('G2', 'TEMP', 'DBEEF4'),
                         ('H2', 'RH', 'DBEEF4'), ('I2', 'TEMP', 'DBEEF4'), ('J2', 'RH', 'DBEEF4'),
                         ('K2', 'TEMP', 'DBEEF4'), ('L2', 'RH', 'DBEEF4'), ('B3', 'Limits →', 'DDD9C3'),
                         ('C3', None, 'F2DCDB'), ('D3', None, 'F2DCDB'), ('E3', None, 'F2DCDB'),
                         ('F3', None, 'F2DCDB'), ('G3', None, 'F2DCDB'), ('H3', None, 'F2DCDB'),
                         ('I3', None, 'F2DCDB'), ('J3', None, 'F2DCDB'), ('K3', None, 'F2DCDB'),
                         ('L3', None, 'F2DCDB')],
               'merges': ['A1:A3', 'C1:D1', 'E1:F1', 'G1:H1', 'I1:J1', 'K1:L1', 'M1:M3'],
               'widths': {'A': 9.14,
                          'B': 12.43,
                          'C': 9.14,
                          'D': 11.0,
                          'E': 9.14,
                          'F': 13.0,
                          'G': 13.0,
                          'H': 13.0,
                          'I': 13.0,
                          'J': 13.0,
                          'K': 13.0,
                          'L': 13.0,
                          'M': 13.0},
               'heights': {1: 25.5}},
              {'name': 'IP_Analytical',
               'first': 9,
               'bcol': 2,
               'ncols': 89,
               'cells': [('A1', None, 'DDD9C3'), ('B1', 'Mini Tablet Batch No. ↓', 'DDD9C3'),
                         ('C1', 'Dry Blend ', 'FF0000'), ('D1', None, '000000'), ('E1', 'Dried Granules', 'FF0000'),
                         ('K1', None, '000000'), ('L1', 'Lubricated Granules', 'DBEEF4'),
                         ('AD1', 'Specification No.', 'DBEEF4'), ('AE1', None, '000000'),
                         ('AF1', 'Uncoated Mini Tablets', 'DBEEF4'), ('AX1', 'Specification No.', 'DBEEF4'),
                         ('AY1', None, '000000'), ('AZ1', 'Film Coated Mini Tablets', 'DBEEF4'),
                         ('BH1', 'Film Coated Mini Tablets', 'DBEEF4'), ('BW1', 'Film Coated Mini Tablets', 'DBEEF4'),
                         ('CJ1', 'In process specification No.', 'CCC1DA'), ('CK1', 'Remarks', 'CCC1DA'),
                         ('A2', 'Sr. No.', 'DDD9C3'), ('B2', 'Test→', 'DDD9C3'),
                         ('C2', 'Loss on Drying (at 105°C) (By Halogen Moisture alyser)', 'DBEEF4'),
                         ('D2', None, '000000'),
                         ('E2', 'Loss on Drying (at 105°C) (By Halogen Moisture alyser) ', 'DBEEF4'),
                         ('F2', 'Bulk density', 'DBEEF4'), ('G2', 'Tapped density', 'DBEEF4'),
                         ('H2', 'Particle size distribution (By Sieve alysis)', 'DBEEF4'), ('K2', None, '000000'),
                         ('L2', 'Assay', 'DBEEF4'),
                         ('M2', 'Loss on Drying (at 105°C) (By Halogen Moisture alyser) ', 'DBEEF4'),
                         ('N2', 'Blend Uniformity', 'DBEEF4'), ('R2', 'Bulk density', 'DBEEF4'),
                         ('S2', 'Tapped density', 'DBEEF4'), ('T2', 'Angle of repose Tanˉ¹  (h/r)', 'DBEEF4'),
                         ('U2', 'Percent compressibility', 'DBEEF4'), ('V2', 'Hausner’s ratio', 'DBEEF4'),
                         ('W2', 'Particle size distribution (By sieve alysis) ', 'DBEEF4'), ('AC2', 'XRD', 'DBEEF4'),
                         ('AE2', None, '000000'), ('AF2', 'Average weight ', 'DBEEF4'),
                         ('AG2', 'Uniformity of dosage units ', 'DBEEF4'), ('AK2', 'Uniformity of mass', 'DBEEF4'),
                         ('AM2', 'Dimensions', 'DBEEF4'), ('AQ2', 'Hardness', 'DBEEF4'),
                         ('AT2', 'Friability', 'DBEEF4'), ('AU2', 'Assay', 'DBEEF4'),
                         ('AV2', 'Loss on drying at 105°C (By Halogen Moisture alyzer)', 'DBEEF4'),
                         ('AW2', 'XRD', 'DBEEF4'), ('AY2', None, '000000'), ('AZ2', 'Average weight ', 'DBEEF4'),
                         ('BA2', 'Uniformity of mass', 'DBEEF4'), ('BC2', 'Uniformity of dosage units ', 'DBEEF4'),
                         ('BG2', 'Loss on drying at 105°C (By Halogen Moisture alyzer)', 'DBEEF4'),
                         ('BH2', 'Dissolution', 'DBEEF4'), ('BW2', 'Assay ', 'DBEEF4'),
                         ('BX2', 'Related Substances by HPLC', 'DBEEF4'),
                         ('CH2', 'Residual Solvents (Ethanol content by GC) ', 'DBEEF4'), ('CI2', 'XRD', 'DBEEF4'),
                         ('D3', None, '000000'), ('K3', None, '000000'), ('N3', 'Avg.', 'DBEEF4'),
                         ('O3', 'Min.', 'DBEEF4'), ('P3', 'Max.', 'DBEEF4'), ('Q3', '%RSD', 'DBEEF4'),
                         ('W3', '% Retained on 30 mesh', 'DBEEF4'), ('X3', '% Retained on 40 mesh', 'DBEEF4'),
                         ('Y3', '% Retained on 60 mesh', 'DBEEF4'), ('Z3', '% Retained on 80 mesh', 'DBEEF4'),
                         ('AA3', '% Retained on 100 mesh', 'DBEEF4'), ('AB3', '% Retained on PAN', 'DBEEF4'),
                         ('AE3', None, '000000'), ('AG3', 'Avg.', 'DBEEF4'), ('AH3', 'Min.', 'DBEEF4'),
                         ('AI3', 'Max.', 'DBEEF4'), ('AJ3', 'Acceptance Value', 'DBEEF4'), ('AY3', None, '000000'),
                         ('BC3', 'Avg.', 'DBEEF4'), ('BD3', 'Min.', 'DBEEF4'), ('BE3', 'Max.', 'DBEEF4'),
                         ('BF3', 'Acceptance Value', 'DBEEF4'), ('BH3', 'Time (hours)', 'DBEEF4'),
                         ('BX3', 'Impurity A', 'DBEEF4'), ('BY3', 'Impurity C', 'DBEEF4'),
                         ('BZ3', 'Impurity D', 'DBEEF4'), ('CA3', 'Impurity E', 'DBEEF4'),
                         ('CB3', 'Impurity F', 'DBEEF4'), ('CC3', 'Impurity  G', 'DBEEF4'),
                         ('CD3', 'Impurity  H', 'DBEEF4'), ('CE3', 'Hydroxy nitrile', 'DBEEF4'),
                         ('CF3', 'Highest Unknown impurity', 'DBEEF4'), ('CG3', 'Total Impurities', 'DBEEF4'),
                         ('D4', None, '000000'), ('K4', None, '000000'), ('AE4', None, '000000'),
                         ('AK4', 'Min Deviation', 'DBEEF4'), ('AL4', 'Max Deviation', 'DBEEF4'),
                         ('AM4', 'Diameter', 'DBEEF4'), ('AO4', 'Thickness', 'DBEEF4'), ('AY4', None, '000000'),
                         ('BA4', 'Min.Deviation', 'DBEEF4'), ('BB4', 'Max.Deviation', 'DBEEF4'), ('BH4', 2, 'DBEEF4'),
                         ('BK4', 4, 'DBEEF4'), ('BN4', 8, 'DBEEF4'), ('BQ4', 12, 'DBEEF4'), ('BT4', 24, 'DBEEF4'),
                         ('D5', None, '000000'), ('K5', None, '000000'), ('AE5', None, '000000'),
                         ('AY5', None, '000000'), ('BH5', '% Drug Release (% of label claim)', 'DBEEF4'),
                         ('D6', None, '000000'), ('K6', None, '000000'), ('AE6', None, '000000'),
                         ('AM6', 'Min. ', 'DBEEF4'), ('AN6', 'Max. ', 'DBEEF4'), ('AO6', 'Min. ', 'DBEEF4'),
                         ('AP6', 'Max. ', 'DBEEF4'), ('AQ6', 'Ave', 'DBEEF4'), ('AR6', 'Min. ', 'DBEEF4'),
                         ('AS6', 'Max. ', 'DBEEF4'), ('AY6', None, '000000'), ('BH6', 'Avg. ', 'DBEEF4'),
                         ('BI6', 'Min. ', 'DBEEF4'), ('BJ6', 'Max. ', 'DBEEF4'), ('BK6', 'Avg. ', 'DBEEF4'),
                         ('BL6', 'Min. ', 'DBEEF4'), ('BM6', 'Max. ', 'DBEEF4'), ('BN6', 'Avg. ', 'DBEEF4'),
                         ('BO6', 'Min. ', 'DBEEF4'), ('BP6', 'Max. ', 'DBEEF4'), ('BQ6', 'Avg. ', 'DBEEF4'),
                         ('BR6', 'Min. ', 'DBEEF4'), ('BS6', 'Max. ', 'DBEEF4'), ('BT6', 'Avg. ', 'DBEEF4'),
                         ('BU6', 'Min. ', 'DBEEF4'), ('BV6', 'Max. ', 'DBEEF4'), ('B7', 'Limits→', 'DDD9C3'),
                         ('C7', 'For information.', 'F2DCDB'), ('D7', None, '000000'), ('E7', 'NMT 3.0', 'F2DCDB'),
                         ('F7', 'For information', 'F2DCDB'), ('G7', 'For information', 'F2DCDB'),
                         ('H7', 'For information.', 'F2DCDB'), ('K7', None, '000000'),
                         ('L7', '95.0 to 105.0 of label claim', 'F2DCDB'), ('M7', 'NMT 3.00', 'F2DCDB'),
                         ('N7', 'NLT 90.0 and NMT 110.0\nRSD NMT 5.0', 'F2DCDB'), ('R7', 'For information', 'F2DCDB'),
                         ('S7', 'For information', 'F2DCDB'), ('T7', 'For information', 'F2DCDB'),
                         ('U7', 'For information', 'F2DCDB'), ('V7', 'For information', 'F2DCDB'),
                         ('W7', 'For information', 'F2DCDB'),
                         ('AC7', 'X-ray diffraction pattern of sample is concordant with   Polymorph Form B',
                          'F2DCDB'),
                         ('AE7', None, '000000'), ('AF7', '22.32 to 24.67', 'F2DCDB'),
                         ('AG7', 'Acceptance value of first 10 dosage units should be less than or equal to 15.0',
                          'F2DCDB'),
                         ('AK7', 'Average mass 10.0 %', 'F2DCDB'), ('AM7', '4.00 ± 0.20 mm', 'F2DCDB'),
                         ('AO7', '1.90 ± 0.40 mm', 'F2DCDB'), ('AQ7', '10 to 60', 'F2DCDB'),
                         ('AT7', 'NMT 0.80', 'F2DCDB'), ('AU7', 'NLT 95.0 and NMT 105.0 of label Claim ', 'F2DCDB'),
                         ('AV7', 'NMT 4.00', 'F2DCDB'),
                         ('AW7', 'X-Ray diffraction pattern of sample is concordant with Polymorph form – B',
                          'F2DCDB'),
                         ('AY7', None, '000000'), ('AZ7', '23.6 to 26.0', 'F2DCDB'),
                         ('BA7', 'Average mass + 10.0%', 'F2DCDB'),
                         ('BC7', 'Acceptance value of first 10 dosage units should be less than or equal to 15.0',
                          'F2DCDB'),
                         ('BG7', 'NMT 4.00', 'F2DCDB'),
                         ('BH7', 'NMT 30/\nBetween 10 to 30/\nBetween 13 to 33', 'F2DCDB'),
                         ('BK7', 'Between 30 to 55/\nBetween 30 to 50/\nBetween 33 to 53', 'F2DCDB'),
                         ('BN7', 'Between 55 to 80/\nBetween 55 to 75/\nBetween 57 to 77', 'F2DCDB'),
                         ('BQ7', 'Between 65 to 90/\nBetween 70 to 90', 'F2DCDB'),
                         ('BT7', 'Not less than 80/\nNot less than 85/\nNot less than 90', 'F2DCDB'),
                         ('BW7', 'NLT 95.0 and NMT 105.0 of label Claim ', 'F2DCDB'), ('BX7', 'NMT 0.10', 'F2DCDB'),
                         ('BY7', 'NMT 0.10', 'F2DCDB'), ('BZ7', 'NMT 0.10', 'F2DCDB'), ('CA7', 'NMT 0.10', 'F2DCDB'),
                         ('CB7', 'NMT 0.10', 'F2DCDB'), ('CC7', 'NMT 0.10', 'F2DCDB'), ('CD7', 'NMT 0.10', 'F2DCDB'),
                         ('CE7', 'NMT 0.10', 'F2DCDB'), ('CF7', 'NMT 0.10', 'F2DCDB'), ('CG7', 'NMT 0.30', 'F2DCDB'),
                         ('CH7', 'NMT 5000/\nNMT 3000', 'F2DCDB'),
                         ('CI7',
                          'X-ray diffraction pattern of sample is concordant with Venlafaxine Hydrochloride '
                          'Polymorph Form B',
                          'F2DCDB'),
                         ('B8', 'Units →', 'DDD9C3'), ('C8', '% w/w', 'F2DCDB'), ('D8', None, '000000'),
                         ('E8', '% w/w', 'F2DCDB'), ('F8', 'g/cc', 'F2DCDB'), ('G8', 'g/cc', 'F2DCDB'),
                         ('H8', '%', 'F2DCDB'), ('K8', None, '000000'), ('L8', '%', 'F2DCDB'), ('M8', '%', 'F2DCDB'),
                         ('N8', '%', 'F2DCDB'), ('R8', 'g/cc', 'F2DCDB'), ('S8', 'g/cc', 'F2DCDB'),
                         ('T8', '%', 'F2DCDB'), ('U8', '%', 'F2DCDB'), ('V8', '%', 'F2DCDB'), ('W8', '%', 'F2DCDB'),
                         ('AA8', None, 'F2DCDB'), ('AB8', None, 'F2DCDB'), ('AC8', None, 'F2DCDB'),
                         ('AE8', None, '000000'), ('AF8', 'mg', 'F2DCDB'), ('AG8', '%', 'F2DCDB'),
                         ('AJ8', None, 'F2DCDB'), ('AK8', None, 'F2DCDB'), ('AL8', None, 'F2DCDB'),
                         ('AM8', 'mm', 'F2DCDB'), ('AO8', 'mm', 'F2DCDB'), ('AQ8', 'N', 'F2DCDB'),
                         ('AT8', '%w/w', 'F2DCDB'), ('AU8', '%', 'F2DCDB'), ('AV8', '%w/w', 'F2DCDB'),
                         ('AW8', None, 'F2DCDB'), ('AY8', None, '000000'), ('AZ8', 'mg', 'F2DCDB'),
                         ('BA8', None, 'F2DCDB'), ('BB8', None, 'F2DCDB'), ('BC8', '%', 'F2DCDB'),
                         ('BF8', None, 'F2DCDB'), ('BG8', '%w/w', 'F2DCDB'), ('BH8', '%', 'F2DCDB'),
                         ('BI8', '%', 'F2DCDB'), ('BJ8', '%', 'F2DCDB'), ('BK8', '%', 'F2DCDB'),
                         ('BL8', '%', 'F2DCDB'), ('BM8', '%', 'F2DCDB'), ('BN8', '%', 'F2DCDB'),
                         ('BO8', '%', 'F2DCDB'), ('BP8', '%', 'F2DCDB'), ('BQ8', '%', 'F2DCDB'),
                         ('BR8', '%', 'F2DCDB'), ('BS8', '%', 'F2DCDB'), ('BT8', '%', 'F2DCDB'),
                         ('BU8', '%', 'F2DCDB'), ('BV8', '%', 'F2DCDB'), ('BW8', '%', 'F2DCDB'),
                         ('BX8', '%', 'F2DCDB'), ('BY8', '%', 'F2DCDB'), ('BZ8', '%', 'F2DCDB'),
                         ('CA8', '%', 'F2DCDB'), ('CB8', '%', 'F2DCDB'), ('CC8', '%', 'F2DCDB'),
                         ('CD8', '%', 'F2DCDB'), ('CE8', '%', 'F2DCDB'), ('CF8', '%', 'F2DCDB'),
                         ('CG8', '%', 'F2DCDB'), ('CH8', 'ppm ', 'F2DCDB'), ('CI8', None, 'F2DCDB')],
               'merges': ['A2:A8', 'AA3:AA6', 'AB3:AB6', 'AC2:AC6', 'AD1:AD8', 'AF1:AW1', 'AF2:AF6', 'AG2:AJ2',
                          'AG3:AG6', 'AG7:AJ7', 'AG8:AI8', 'AH3:AH6', 'AI3:AI6', 'AJ3:AJ6', 'AK2:AL3', 'AK4:AK6',
                          'AK7:AL7', 'AL4:AL6', 'AM2:AP3', 'AM4:AN5', 'AM7:AN7', 'AM8:AN8', 'AO4:AP5', 'AO7:AP7',
                          'AO8:AP8', 'AQ2:AS5', 'AQ7:AS7', 'AQ8:AS8', 'AT2:AT6', 'AU2:AU6', 'AV2:AV6', 'AW2:AW6',
                          'AX1:AX8', 'AZ1:BG1', 'AZ2:AZ6', 'B2:B6', 'BA2:BB3', 'BA4:BA6', 'BA7:BB7', 'BB4:BB6',
                          'BC2:BF2', 'BC3:BC6', 'BC7:BF7', 'BC8:BE8', 'BD3:BD6', 'BE3:BE6', 'BF3:BF6', 'BG2:BG6',
                          'BH1:BV1', 'BH2:BV2', 'BH3:BV3', 'BH4:BJ4', 'BH5:BV5', 'BH7:BJ7', 'BK4:BM4', 'BK7:BM7',
                          'BN4:BP4', 'BN7:BP7', 'BQ4:BS4', 'BQ7:BS7', 'BT4:BV4', 'BT7:BV7', 'BW1:CI1', 'BW2:BW6',
                          'BX2:CG2', 'BX3:BX6', 'BY3:BY6', 'BZ3:BZ6', 'C2:C6', 'CA3:CA6', 'CB3:CB6', 'CC3:CC6',
                          'CD3:CD6', 'CE3:CE6', 'CF3:CF6', 'CG3:CG6', 'CH2:CH6', 'CI2:CI6', 'CJ1:CJ8', 'CK1:CK8',
                          'E1:J1', 'E2:E6', 'F2:F6', 'G2:G6', 'H2:J6', 'H7:J7', 'H8:J8', 'L1:AC1', 'L2:L6', 'M2:M6',
                          'N2:Q2', 'N3:N6', 'N7:Q7', 'N8:Q8', 'O3:O6', 'P3:P6', 'Q3:Q6', 'R2:R6', 'S2:S6', 'T2:T6',
                          'U2:U6', 'V2:V6', 'W2:AB2', 'W3:W6', 'W7:AB7', 'W8:Z8', 'X3:X6', 'Y3:Y6', 'Z3:Z6'],
               'widths': {'A': 9.14,
                          'B': 10.29,
                          'C': 9.14,
                          'D': 0.43,
                          'E': 9.14,
                          'F': 13.0,
                          'G': 13.0,
                          'I': 13.0,
                          'J': 13.0,
                          'K': 0.86,
                          'L': 9.14,
                          'M': 13.0,
                          'N': 13.0,
                          'O': 13.0,
                          'P': 13.0,
                          'Q': 13.0,
                          'R': 13.0,
                          'S': 13.0,
                          'T': 13.0,
                          'U': 13.0,
                          'V': 13.0,
                          'W': 13.0,
                          'X': 13.0,
                          'Y': 13.0,
                          'Z': 13.0,
                          'AA': 13.0,
                          'AB': 13.0,
                          'AC': 13.0,
                          'AD': 24.57,
                          'AE': 3.86,
                          'AF': 9.14,
                          'AG': 13.0,
                          'AH': 13.0,
                          'AI': 13.0,
                          'AJ': 13.0,
                          'AK': 13.0,
                          'AL': 13.0,
                          'AM': 13.0,
                          'AN': 13.0,
                          'AO': 13.0,
                          'AP': 13.0,
                          'AQ': 13.0,
                          'AR': 13.0,
                          'AS': 13.0,
                          'AT': 13.0,
                          'AU': 13.0,
                          'AV': 13.0,
                          'AW': 9.0,
                          'AX': 24.0,
                          'AY': 4.14,
                          'AZ': 9.14,
                          'BA': 10.14,
                          'BB': 10.43,
                          'BC': 9.14,
                          'BD': 13.0,
                          'BE': 13.0,
                          'BF': 13.0,
                          'BG': 13.0,
                          'BH': 13.0,
                          'BI': 13.0,
                          'BJ': 13.0,
                          'BK': 13.0,
                          'BL': 13.0,
                          'BM': 13.0,
                          'BN': 13.0,
                          'BO': 13.0,
                          'BP': 13.0,
                          'BQ': 13.0,
                          'BR': 13.0,
                          'BS': 13.0,
                          'BT': 13.0,
                          'BU': 13.0,
                          'BV': 13.0,
                          'BW': 13.0,
                          'BX': 13.0,
                          'BY': 13.0,
                          'BZ': 13.0,
                          'CA': 13.0,
                          'CB': 13.0,
                          'CC': 13.0,
                          'CD': 13.0,
                          'CE': 13.0,
                          'CF': 13.0,
                          'CG': 13.0,
                          'CH': 13.0,
                          'CI': 13.0,
                          'CJ': 20.71,
                          'CK': 23.71},
               'heights': {1: 37.5, 2: 15.0, 4: 15.0, 7: 47.25}}]}



if __name__ == '__main__':
    main()
